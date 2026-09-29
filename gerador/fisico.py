"""Curvas de avanço físico por cenário, a partir das PLs (itens DIRETOS), para o painel CURVA FÍSICA.

Uso: python3 fisico.py saida.json "Padrão=arq_calc.xlsx" "Prevision=arq_calc.xlsx" ...

Por PL (aba de disciplina):
  realizado (até o corte)   = % realizado acumulado do Prevision (coluna O)
  previsto Prevision        = % previsto acumulado do Prevision (coluna M)
  cenário (depois do corte) = físico no corte + (1 − físico no corte) × desembolso acumulado depois do corte ÷ desembolso total
                              depois do corte (colunas C + D do cenário). PL sem desembolso futuro segue o previsto do Prevision.
Obra = média das PLs diretas ponderada pela projeção da PL (total no fluxo, R$).
"""
import sys, json, datetime as dt
from openpyxl import load_workbook

R0, R1, RTOT = 5, 76, 78


def num(v): return float(v) if isinstance(v, (int, float)) else 0.0


def ler(arq):
    wb = load_workbook(arq, data_only=True, read_only=False)
    vp = wb['VALOR POR PL']; corte = vp['P5'].value
    tabs = [n for n in wb.sheetnames if wb[n]['I5'].value == "VALOR ORÇADO (R$ base / INCC)" and wb[n]['K13'].value == 'DIRETO']
    meses = [wb[tabs[0]].cell(r, 2).value for r in range(R0, R1 + 1)]
    pls = []
    for n in tabs:
        ws = wb[n]
        w = num(ws.cell(RTOT, 3).value) + num(ws.cell(RTOT, 4).value)
        if w <= 0: continue
        prev = [min(1, num(ws.cell(r, 13).value) / 100) for r in range(R0, R1 + 1)]
        real = [min(1, num(ws.cell(r, 15).value) / 100) if meses[i] <= corte else None for i, r in enumerate(range(R0, R1 + 1))]
        fc = min(1, num(ws['K14'].value))
        fut = [(num(ws.cell(r, 3).value) + num(ws.cell(r, 4).value)) if meses[i] > corte else 0 for i, r in enumerate(range(R0, R1 + 1))]
        tf = sum(fut); acc = 0; cen = []
        for i in range(len(meses)):
            if meses[i] <= corte: cen.append(real[i]); continue
            acc += fut[i]
            cen.append(fc + (1 - fc) * acc / tf if tf > 0 else max(fc, prev[i]))
        pls.append(dict(n=ws['B2'].value or n, aba=n, w=w, prev=prev, real=real, cen=cen, fc=fc))
    return corte, meses, pls


def pond(pls, key, i):
    sw = sum(p['w'] for p in pls)
    return sum(p['w'] * (p[key][i] or 0) for p in pls) / sw if sw else 0


def main():
    out, pares = sys.argv[1], [a.split('=', 1) for a in sys.argv[2:]]
    cens, base = [], None
    for nome, arq in pares:
        corte, meses, pls = ler(arq)
        if base is None: base = (corte, meses, pls)
        cens.append(dict(nome=nome, v=[round(pond(pls, 'cen', i), 6) for i in range(len(meses))],
                         pls=[dict(n=p['n'], w=round(p['w'], 2), fc=round(p['fc'], 4), pdez=[round(p['prev'][i], 4) for i in range(len(meses)) if meses[i].month == 12], dez=[round(p['cen'][i], 4) for i in range(len(meses)) if meses[i].month == 12])
                              for p in pls]))
    corte, meses, pls = base
    ic = max(i for i, m in enumerate(meses) if m <= corte)
    dados = dict(corte=corte.strftime('%Y-%m'), meses=[m.strftime('%Y-%m') for m in meses],
                 real=[round(pond(pls, 'real', i), 6) if i <= ic else None for i in range(len(meses))],
                 prev=[round(pond(pls, 'prev', i), 6) for i in range(len(meses))],
                 cens=cens, npl=len(pls),
                 anos=[m.year for m in meses if m.month == 12])
    json.dump(dados, open(out, 'w'), ensure_ascii=False, separators=(',', ':'))
    print('corte', dados['corte'], 'realizado', round(dados['real'][ic] * 100, 2), '% · previsto Prevision no corte', round(dados['prev'][ic] * 100, 2), '%')
    for c in cens:
        dz = {m: round(v * 100, 1) for m, v in zip(dados['meses'], c['v']) if m.endswith('-12')}
        print(c['nome'], dz)
    print('Prevision', {m: round(v * 100, 1) for m, v in zip(dados['meses'], dados['prev']) if m.endswith('-12')})


if __name__ == '__main__':
    main()
