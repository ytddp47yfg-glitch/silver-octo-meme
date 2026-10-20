"""Carrega o export de ATIVIDADES do Prevision (cronograma) na chave CRONO do data2.json das pastas indicadas.

Uso: python3 -I gerador/prevision/crono.py <export.xlsx> <obra> <pasta> [<pasta> ...]
     ex.: ... crono.py cronograma_ed-jardim_19-10.xlsx ed-jardim atual
O export tem o cabeçalho nas linhas 1–3 e uma atividade por linha a partir da 4 (colunas B:K = ID, pacote, serviço,
lote, início, término, data de referência, base, previsto, realizado). Faz backup (data2_antes_crono.json, fora do git)
e mostra o que mudou: nº de atividades, data de referência e atividades com término alterado.
"""
import datetime as dt, json, os, shutil, sys
from openpyxl import load_workbook

arq, obra, pastas = sys.argv[1], sys.argv[2], sys.argv[3:]
if not pastas: sys.exit(__doc__)
F = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "ferramenta", "obras", obra)
rows = list(load_workbook(arq, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
iso = lambda v: v.isoformat() if isinstance(v, dt.datetime) else v
for p in pastas:
    d2p = os.path.join(F, p, "data2.json"); D = json.load(open(d2p)); old = D["CRONO"]
    new = {"A1": old.get("A1") or rows[0][0], "A2": rows[1][0] or old.get("A2"), "I3": old.get("I3")}   # A1 = título da obra (fica)
    for i, r in enumerate(rows[3:], start=4):
        if i > 4 and r[1] in ("", None): continue
        for j, c in enumerate("BCDEFGHIJK", start=1):
            if j < len(r) and r[j] not in ("", None): new[f"{c}{i}"] = iso(r[j])
    fim = lambda C: {C.get(f"B{k[1:]}"): C[k] for k in C if k[0] == "G" and k[1:].isdigit()}
    fo, fn = fim(old), fim(new); mud = [a for a in fn if a in fo and fo[a] != fn[a]]
    print(f"{obra}/{p}: atividades {sum(k[0] == 'B' for k in old)} → {sum(k[0] == 'B' for k in new)} | "
          f"referência {old.get('H5')} → {new.get('H5')} | término alterado em {len(mud)} | novas {len(set(fn) - set(fo))} | saíram {len(set(fo) - set(fn))}")
    bk = os.path.join(F, p, "data2_antes_crono.json")
    if not os.path.exists(bk): shutil.copy(d2p, bk)
    D["CRONO"] = new; json.dump(D, open(d2p, "w"), ensure_ascii=False)
print("Agora: python3 ferramenta/regerar.py", obra, "(a auditoria acusa pacote novo sem de-para pacote → PL)")
