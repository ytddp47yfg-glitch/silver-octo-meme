"""Carrega o CRONOGRAMA FÍSICO-FINANCEIRO (CFF) exportado do Prevision na chave CURVA PREVISION das pastas indicadas.

Uso: python3 -I gerador/prevision/cff.py <cff.xlsx> <obra> <pasta> [<pasta> ...]
Copia a 1ª aba célula a célula (mesmo layout de sempre), compara com a curva atual (itens EAP, última coluna/mês) e faz
backup (data2_antes_cff.json, fora do git). Atenção: cenários com curva ajustada à mão (ex.: Ed. Jardim "atual", cuja
curva foi deslocada pelas datas do cronograma de atividades) perdem o ajuste — refaça-o ou não inclua a pasta.
"""
import json, os, re, shutil, sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as CI, get_column_letter as L

arq, obra, pastas = sys.argv[1], sys.argv[2], sys.argv[3:]
if not pastas: sys.exit(__doc__)
F = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "ferramenta", "obras", obra)
ws = load_workbook(arq, data_only=True).worksheets[0]
new = {c.coordinate: (c.value.isoformat() if hasattr(c.value, "isoformat") else c.value)
       for row in ws.iter_rows() for c in row if c.value not in (None, "")}
cod = lambda d: [v for _, v in sorted((int(k[1:]), v) for k, v in d.items() if re.fullmatch(r"B\d+", k))]
ult = lambda d: max((CI(re.match("[A-Z]+", k)[0]) for k in d if re.fullmatch(r"[A-Z]+2", k)), default=0)
for p in pastas:
    d2p = os.path.join(F, p, "data2.json"); D = json.load(open(d2p)); old = D["CURVA PREVISION"]
    print(f"{obra}/{p}: itens EAP {len(cod(old))} → {len(cod(new))} (mesma ordem: {cod(old) == cod(new)}) | "
          f"última coluna {L(ult(old))} → {L(ult(new))} | A1: {new.get('A1')}")
    bk = os.path.join(F, p, "data2_antes_cff.json")
    if not os.path.exists(bk): shutil.copy(d2p, bk)
    D["CURVA PREVISION"] = new; json.dump(D, open(d2p, "w"), ensure_ascii=False)
print("Agora: python3 ferramenta/regerar.py", obra)
