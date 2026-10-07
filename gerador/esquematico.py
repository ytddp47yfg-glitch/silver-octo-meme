"""Aba ESQUEMÁTICO: corte ilustrativo do prédio com o andamento de cada PL por pavimento.

esquematico(wb, obra): cria a aba a partir da aba CRONOGRAMA ATIVIDADES da ferramenta (export do Prevision já com
o de-para pacote → PL na coluna M e o nº do pavimento na coluna N). Tudo é fórmula: colar um cronograma novo na aba
de atividades e recalcular atualiza o esquemático.

Dois cortes lado a lado:
  1. situação na data de referência do cronograma (% realizado do Prevision, coluna K);
  2. previsão para uma data digitável (padrão = data do RETRATO): atividades com saldo seguem linearmente
     da data de referência (ou do início, se depois) até o término planejado.
Célula = avanço do PL no pavimento, média das tarefas ponderada pela duração (linhas-resumo '-' só entram quando
o pacote/lote não tem tarefas). Base: só pavimentos numerados (coluna N); infraestrutura e fachada em faixas próprias.

Uso como módulo (build da ferramenta): import esquematico; esquematico.esquematico(wb, "ED JARDIM")
Uso avulso: python3 esquematico.py ferramenta.xlsx saida.xlsx   (relê e grava a pasta com openpyxl; prefira o build)
"""
import re, sys, unicodedata
from collections import Counter, defaultdict
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

F_ = "Calibri"
NAVY, INK2, BG, CARD, LINE = "1C3B2A", "8C877A", "F4F1EA", "FCFBF8", "E5DFD2"
C_OK, C_AND, C_AND2, C_NAO, C_SEM = "1F7A52", "C0A062", "E8D9B0", "D9D4C7", "F4F1EA"
fill = lambda h: PatternFill("solid", fgColor=h)
fino = Side(style="thin", color="FCFBF8")
CA = "'CRONOGRAMA ATIVIDADES'!"
R0A, R1A = 5, 6004            # linhas de atividades da aba CRONOGRAMA ATIVIDADES
HC = 60                      # 1ª coluna auxiliar (oculta) nesta aba
MIN_PAV = 3                  # PL entra no corte se tiver atividades em pelo menos 3 pavimentos


def _norm(s): return unicodedata.normalize("NFKD", str(s).upper()).encode("ascii", "ignore").decode()


def _pav(lote):
    """Mesma regra da coluna N da aba CRONOGRAMA ATIVIDADES."""
    if not lote: return None
    s = str(lote)
    if "SUBSOLO" in s.upper():
        m = re.match(r"\s*(\d+)\s*[º°]", s); return -int(m.group(1)) if m else -1
    if s.startswith("TÉRREO"): return 0
    m = re.match(r"\s*(\d+)\s*[º°]", s)
    return int(m.group(1)) if m else None


def _rotulo(pav, lotes):
    if pav == 0: return "TÉRREO"
    nome = Counter(lotes).most_common(1)[0][0]
    nome = re.sub(r"^\s*\d+\s*[º°]\s*(PAV\.?)?\s*[-–]?\s*", "", str(nome)).strip(" .-")
    nome = re.sub(r"\s+", " ", nome)
    return f"{pav}º  {nome[:22]}" if pav > 0 else f"{-pav}º SUBSOLO"


def _ler(wb):
    ca, vp = wb["CRONOGRAMA ATIVIDADES"], wb["VALOR POR PL"]
    depara = {str(ca.cell(r, 18).value).strip(): str(ca.cell(r, 19).value).strip()
              for r in range(5, 305) if ca.cell(r, 18).value and ca.cell(r, 19).value not in (None, "")}
    nome_pl = {str(vp.cell(r, 3).value).strip(): str(vp.cell(r, 4).value).strip() for r in range(11, 49) if vp.cell(r, 3).value}
    pav_pl, lotes = defaultdict(set), defaultdict(list)
    sem_pav = Counter()
    for r in range(R0A, R1A + 1):
        pac = ca.cell(r, 3).value
        if not pac: continue
        pl = depara.get(str(pac).strip())
        p = _pav(ca.cell(r, 5).value)
        if p is None:
            if pl: sem_pav[pl] += 1
            continue
        lotes[p].append(ca.cell(r, 5).value)
        if pl: pav_pl[pl].add(p)
    pls = sorted((pl for pl, s in pav_pl.items() if len(s) >= MIN_PAV), key=lambda x: (len(x), x))
    pavs = sorted(lotes, reverse=True)
    infra = [c for c, n in nome_pl.items() if re.search(r"CONTEN|TERRAPL|FUNDA", _norm(n))]
    fach = [c for c, n in nome_pl.items() if "FACHADA" in _norm(n)]
    return pls, pavs, {p: _rotulo(p, lotes[p]) for p in pavs}, nome_pl, infra, fach


def _curto(n):
    n = _norm(n)
    for a, b in (("INSTALACOES", "INST."), ("REVESTIMENTOS", "REV."), ("REVESTIMENTO", "REV."), ("COMBATE ", ""),
                 ("DE MASSA EM ", "MASSA "), ("ESTRUTURA EM CONCRETO ARMADO", "ESTRUTURA"), ("DE ACABAMENTO", "ACABAM."),
                 ("HIDROSANITARIAS", "HIDROSSANIT."), ("IMPERMEABILIZACAO", "IMPERMEABILIZ."), ("ESPECIAIS", "ESPECIAIS"),
                 (" DIVERSOS", ""), (" E METAIS", "/METAIS"), ("LIMPEZA FINAL DE OBRA", "LIMPEZA FINAL")):
        n = n.replace(a, b)
    return n.strip()


def esquematico(wb, obra):
    if "CRONOGRAMA ATIVIDADES" not in wb.sheetnames: return None
    pls, pavs, rot, nome_pl, infra, fach = _ler(wb)
    if not pls or not pavs: return None
    pos = max(wb.sheetnames.index(n) for n in ("RETRATO 2027", "CURVA FÍSICA", "GANTT") if n in wb.sheetnames) + 1
    ws = wb.create_sheet("ESQUEMÁTICO", pos)
    ws.sheet_view.showGridLines = False; ws.sheet_properties.tabColor = "1F7A52"
    ws.sheet_view.zoomScale = 90
    NP = len(pls)
    # colunas: A (nº pav, oculto) | B rótulo | PLs | fachada | geral | espaço | rótulo | PLs | fachada | geral
    B1 = 3; FA1 = B1 + NP; GE1 = FA1 + 1
    LB2 = GE1 + 3; B2 = LB2 + 1; FA2 = B2 + NP; GE2 = FA2 + 1
    LAST = GE2 + 1
    ws.column_dimensions["A"].width = 3; ws.column_dimensions["A"].hidden = True
    for c in (2, LB2): ws.column_dimensions[L(c)].width = 24
    for b in (B1, B2):
        for k in range(NP): ws.column_dimensions[L(b + k)].width = 3.6
    for c in (FA1, FA2): ws.column_dimensions[L(c)].width = 4.2
    for c in (GE1, GE2): ws.column_dimensions[L(c)].width = 11
    for c in (GE1 + 1, GE1 + 2): ws.column_dimensions[L(c)].width = 3
    ws.column_dimensions[L(LAST)].width = 3

    ws["B2"] = f"{obra}  ·  ESQUEMÁTICO DE AVANÇO POR PAVIMENTO"; ws["B2"].font = Font(name=F_, bold=True, size=20, color=NAVY)
    ws["B3"] = ("Corte ilustrativo (sem escala): cada célula é o avanço da PL no pavimento, pelas atividades da aba "
                "CRONOGRAMA ATIVIDADES (média das tarefas ponderada pela duração).")
    ws["B3"].font = Font(name=F_, size=10, color=INK2)
    ws.row_dimensions[2].height = 34; ws.row_dimensions[3].height = 20

    # datas
    ws["B5"] = "Situação do cronograma em"; ws["B6"] = "Previsão para (digite a data)"
    for r in (5, 6): ws[f"B{r}"].font = Font(name=F_, bold=True, size=10, color=NAVY)
    ws["C5"] = f"=MAX({CA}$H${R0A}:$H${R1A})"
    ws["C6"] = "='RETRATO 2027'!$K$5" if "RETRATO 2027" in wb.sheetnames else "=DATE(YEAR(C5),12,31)"
    for r, f_ in ((5, Font(name=F_, bold=True, size=11, color=NAVY)), (6, Font(name=F_, bold=True, size=11, color="0000FF"))):
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=3 + 5)
        c = ws.cell(r, 3); c.font = f_; c.number_format = "dd/mm/yyyy"; c.alignment = Alignment(horizontal="left")
    ws["C6"].fill = fill("FFF2CC")
    ws["C6"].comment = Comment("Data do segundo corte. Por padrão é a data do RETRATO (aba RETRATO 2027, K5); pode digitar "
                               "qualquer data posterior à situação do cronograma. Antes dela vale o realizado.", "Ferramenta")
    # legenda
    leg = [(C_OK, "Concluído (100%)"), (C_AND, "Em andamento (50% a 99%)"), (C_AND2, "Em andamento (1% a 49%)"),
           (C_NAO, "A executar (0%)"), (C_SEM, "Sem atividade da PL no pavimento")]
    for i, (cor, t) in enumerate(leg):
        r = 4 + i; c = B2 + 1
        cl = ws.cell(r, c); cl.fill = fill(cor); cl.border = Border(left=Side(style="thin", color=LINE), right=Side(style="thin", color=LINE),
                                                                    top=Side(style="thin", color=LINE), bottom=Side(style="thin", color=LINE))
        ws.cell(r, c + 1, t).font = Font(name=F_, size=9, color=INK2)

    # colunas auxiliares (ocultas): peso = duração das tarefas; % na data do 2º corte
    W, P = L(HC), L(HC + 1)
    ws.cell(4, HC, "peso (dias)").font = Font(name=F_, size=8, color=INK2)
    ws.cell(4, HC + 1, "% na data C6").font = Font(name=F_, size=8, color=INK2)
    for r in range(R0A, R1A + 1):
        a = lambda col: f"{CA}{col}{r}"
        ws.cell(r, HC, f'=IF(OR({a("C")}="",NOT(ISNUMBER({a("F")})),NOT(ISNUMBER({a("G")}))),0,'
                       f'IF({a("D")}="-",IF(COUNTIFS({CA}$C${R0A}:$C${R1A},{a("C")},{CA}$E${R0A}:$E${R1A},{a("E")},{CA}$D${R0A}:$D${R1A},"<>-")>0,0,'
                       f'MAX(1,{a("G")}-{a("F")})),MAX(1,{a("G")}-{a("F")})))')
        ws.cell(r, HC + 1, f'=IF({W}{r}=0,0,IF($C$6<={a("H")},{a("K")},{a("K")}+(100-{a("K")})*MIN(1,MAX(0,($C$6-MAX({a("F")},{a("H")}))'
                           f'/MAX(1,{a("G")}-MAX({a("F")},{a("H")}))))))')
    for c in (HC, HC + 1): ws.column_dimensions[L(c)].hidden = True
    WR = f"${W}${R0A}:${W}${R1A}"
    PR = {1: f"{CA}$K${R0A}:$K${R1A}", 2: f"${P}${R0A}:${P}${R1A}"}
    MR, NR = f"{CA}$M${R0A}:$M${R1A}", f"{CA}$N${R0A}:$N${R1A}"

    def media(k, cond):
        return f"=IFERROR(SUMPRODUCT({WR},{PR[k]},{cond})/SUMPRODUCT({WR},{cond}),\"\")"

    def em(codes): return "(" + "+".join(f'({MR}="{c}")' for c in codes) + ">0)*1"

    HT, HN, RP0 = 9, 10, 12          # título do bloco, nomes das PLs, 1º pavimento (linha 11 = cobertura)
    RPN = RP0 + len(pavs) - 1
    RINF, RTER, RTOT = RPN + 1, RPN + 2, RPN + 4
    ws.row_dimensions[HN].height = 118
    for r in range(RP0, RPN + 1): ws.row_dimensions[r].height = 14
    ws.row_dimensions[RP0 - 1].height = 6; ws.row_dimensions[RINF].height = 20; ws.row_dimensions[RTER].height = 10
    titulos = {1: ('="SITUAÇÃO EM "&TEXT($C$5,"dd/mm/yyyy")&"  (realizado do cronograma)"'),
               2: ('="PREVISÃO PARA "&TEXT($C$6,"dd/mm/yyyy")&"  (cronograma planejado)"')}
    for k, (lb, b, fa, ge) in {1: (2, B1, FA1, GE1), 2: (LB2, B2, FA2, GE2)}.items():
        c = ws.cell(HT, lb, titulos[k]); c.font = Font(name=F_, bold=True, size=12, color=NAVY)
        hd = Font(name=F_, bold=True, size=8, color=NAVY)
        for j, pl in enumerate(pls):
            c = ws.cell(HN, b + j, _curto(nome_pl.get(pl, pl))); c.font = hd
            c.alignment = Alignment(text_rotation=90, horizontal="center", vertical="bottom")
        for col, t in ((fa, "FACHADA"), (ge, "PAVIMENTO (todas as PLs)")):
            c = ws.cell(HN, col, t); c.font = hd
            c.alignment = Alignment(text_rotation=90 if col == fa else 0, horizontal="center", vertical="bottom", wrap_text=True)
        # cobertura
        for c_ in range(b - 1, fa + 1): ws.cell(RP0 - 1, c_).fill = fill(NAVY)
        # pavimentos
        for i, pv in enumerate(pavs):
            r = RP0 + i
            ws.cell(r, 1, pv).font = Font(name=F_, size=8, color=INK2)
            c = ws.cell(r, lb, rot[pv]); c.font = Font(name=F_, size=8, bold=pv in (0, max(pavs)), color=NAVY)
            c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
            for j, pl in enumerate(pls):
                c = ws.cell(r, b + j, media(k, f'({MR}="{pl}")*({NR}=$A{r})'))
                c.number_format = '0;-0;0;@'; c.font = Font(name=F_, size=7, color="3C3A33")
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.fill = fill(C_SEM); c.border = Border(left=fino, right=fino, top=fino, bottom=fino)
            g = ws.cell(r, ge, media(k, f'({MR}<>"")*({NR}=$A{r})')); g.number_format = '0"%"'
            g.font = Font(name=F_, size=8, bold=True, color=NAVY); g.alignment = Alignment(horizontal="right")
        # fachada: faixa vertical ao lado do corte
        ws.merge_cells(start_row=RP0, start_column=fa, end_row=RPN, end_column=fa)
        c = ws.cell(RP0, fa, media(k, em(fach)) if fach else '=""'); c.number_format = '0"%"'
        c.font = Font(name=F_, bold=True, size=9, color=NAVY); c.alignment = Alignment(text_rotation=90, horizontal="center", vertical="center")
        c.fill = fill(C_SEM)
        # infraestrutura (contenção, terraplenagem, fundações) e terreno
        ws.cell(RINF, 1).value = None
        ws.merge_cells(start_row=RINF, start_column=b, end_row=RINF, end_column=fa)
        vi = media(k, em(infra))[1:] if infra else '""'
        c = ws.cell(RINF, b, f'=IFERROR("INFRAESTRUTURA (contenção, terraplenagem e fundações): "&TEXT({vi},"0")&"%","")')
        c.font = Font(name=F_, bold=True, size=8, color=NAVY); c.alignment = Alignment(horizontal="center", vertical="center")
        c.fill = fill(C_SEM)
        ws.cell(RINF, ge, f"={vi}" if infra else '=""').number_format = '0"%"'
        ws.cell(RINF, ge).font = Font(name=F_, size=8, bold=True, color=NAVY)
        for c_ in range(b - 1, ge + 1): ws.cell(RTER, c_).fill = PatternFill("lightUp", fgColor="8C877A", bgColor="E5DFD2")
        # % de cada PL em todos os pavimentos
        c = ws.cell(RTOT, lb, "% da PL nos pavimentos"); c.font = Font(name=F_, bold=True, size=8, color=NAVY)
        c.alignment = Alignment(horizontal="right", indent=1)
        for j, pl in enumerate(pls):
            c = ws.cell(RTOT, b + j, media(k, f'({MR}="{pl}")*ISNUMBER({NR})')); c.number_format = '0'
            c.font = Font(name=F_, size=7, bold=True, color=NAVY); c.alignment = Alignment(horizontal="center")
        g = ws.cell(RTOT, ge, media(k, f'({MR}<>"")*ISNUMBER({NR})')); g.number_format = '0"%"'
        g.font = Font(name=F_, size=9, bold=True, color=NAVY)
        # cores por situação
        rng = f"{L(b)}{RP0}:{L(fa)}{RPN}"; tl = f"{L(b)}{RP0}"
        regras = [(f'AND(ISNUMBER({tl}),{tl}>=99.95)', C_OK, "FFFFFF"), (f'AND(ISNUMBER({tl}),{tl}>=50)', C_AND, "3C3A33"),
                  (f'AND(ISNUMBER({tl}),{tl}>0.05)', C_AND2, "3C3A33"), (f'ISNUMBER({tl})', C_NAO, C_NAO)]
        for fml, cor, fc in regras:
            ws.conditional_formatting.add(rng, FormulaRule(formula=[fml], fill=fill(cor), font=Font(name=F_, color=fc), stopIfTrue=True))
        # infraestrutura: cor pelo valor da coluna "geral"
        rinf = f"{L(b)}{RINF}:{L(fa)}{RINF}"; vref = f"${L(ge)}${RINF}"
        for fml, cor, fc in [(f'AND(ISNUMBER({vref}),{vref}>=99.95)', C_OK, "FFFFFF"), (f'AND(ISNUMBER({vref}),{vref}>=50)', C_AND, NAVY),
                             (f'AND(ISNUMBER({vref}),{vref}>0.05)', C_AND2, NAVY), (f'ISNUMBER({vref})', C_NAO, NAVY)]:
            ws.conditional_formatting.add(rinf, FormulaRule(formula=[fml], fill=fill(cor), font=Font(name=F_, bold=True, color=fc), stopIfTrue=True))
        ws.conditional_formatting.add(f"{L(ge)}{RP0}:{L(ge)}{RPN}",
                                      DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color=C_OK, showValue=True))
    nota = (f"Pavimentos: nº na coluna N da aba CRONOGRAMA ATIVIDADES (lote do Prevision). PLs com atividades em pelo menos "
            f"{MIN_PAV} pavimentos; infraestrutura = {', '.join(nome_pl[c].lower() for c in infra) or '—'}; "
            f"fachada = {', '.join(nome_pl[c].lower() for c in fach) or '—'} (lotes sem pavimento). "
            "Previsão: o saldo de cada tarefa segue linear da data da situação (ou do início planejado) até o término planejado. "
            "Esquema ilustrativo, sem escala, não substitui o levantamento em campo.")
    ws.merge_cells(start_row=RTOT + 2, start_column=2, end_row=RTOT + 2, end_column=LAST)
    c = ws.cell(RTOT + 2, 2, nota); c.font = Font(name=F_, italic=True, size=8, color=INK2)
    c.alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[RTOT + 2].height = 26
    # fundo creme no que ficou sem preenchimento
    for row in ws.iter_rows(min_row=1, max_row=RTOT + 3, min_col=1, max_col=LAST):
        for c in row:
            if type(c).__name__ != "MergedCell" and (c.fill is None or c.fill.fill_type is None): c.fill = fill(BG)
    ws.freeze_panes = None
    ws.sheet_view.selection[0].activeCell = "C6"; ws.sheet_view.selection[0].sqref = "C6"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{L(LAST)}{RTOT + 2}"
    return dict(pls=pls, pavs=pavs, infra=infra, fachada=fach)


if __name__ == "__main__":
    from openpyxl import load_workbook
    wb = load_workbook(sys.argv[1])
    if "ESQUEMÁTICO" in wb.sheetnames: del wb["ESQUEMÁTICO"]
    vp = wb["VALOR POR PL"]
    print(esquematico(wb, str(vp["P9"].value or "").strip()))
    wb.save(sys.argv[2])
