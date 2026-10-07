"""Aba AUX.ESTRUTURA: quantitativos da estrutura por trecho/pavimento → curva mensal da aba ESTRUTURA.

Resgata a aba auxiliar da planilha de origem (Torre Catharina: AUX.STR.1; Botânico: AUX.ESTRUTURA) com as fórmulas
internas. Referências externas ([18]Orçamento, [19]Quantitativos…) e a outras abas entram como valor. As datas de cada
trecho passam a vir da aba CRONOGRAMA ATIVIDADES (linha-resumo do pacote/lote no Prevision); sem par no cronograma,
vale a data da origem. A aba ESTRUTURA usa a curva resultante como PREVISÃO INFORMADA (coluna H): o saldo da meta
em INCC é distribuído nos meses futuros na proporção dela.

  tipo "trechos" (Torre Catharina): componentes por mês de TÉRMINO do trecho, regras de pagamento da origem
      (retenção técnica sobre forma e aço MDO; aço material 88% em 2 parcelas + 12% corte e dobra) e pagamento
      no mês seguinte ao custo.
  tipo "pavimentos" (Botânico): custo de cada pavimento rateado pelos dias de execução (início → término) em cada
      mês; pavimento sem par no cronograma segue a distribuição manual da origem (colunas AC:AX).

Extrair da origem (uma vez):  python3 aux_estrutura.py extrair origem.xlsx "AUX.STR.1" trechos  saida.json
Build:  cfg.json da obra → "aux_estrutura": "aux_estrutura.json" (arquivo na pasta de dados); build_modelo.py chama
        montar(wb, A, ...) e liga ESTRUTURA!H à curva.
"""
import json, re, sys, datetime as dt
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

F_ = "Calibri"
NAVY, INK2, BG = "1C3B2A", "8C877A", "F4F1EA"
fill = lambda h: PatternFill("solid", fgColor=h)
CA = "'CRONOGRAMA ATIVIDADES'!"
RA = "$5:$"  # placeholder
NOME = "AUX.ESTRUTURA"


# ------------------------------------------------------------------ extração (planilha de origem → JSON)
def extrair(arq, aba, tipo, saida):
    from openpyxl import load_workbook
    wf, wv = load_workbook(arq)[aba], load_workbook(arq, data_only=True)[aba]
    proprio = [f"'{aba}'!", f"{aba}!"]
    cells = {}
    for row in wf.iter_rows():
        for c in row:
            f, v = c.value, wv[c.coordinate].value
            if hasattr(f, "text"): f = f.text
            if f is None and v is None: continue
            if isinstance(v, (dt.datetime, dt.date)): v = v.isoformat()
            if isinstance(f, str) and f.startswith("="):
                g = f
                for p in proprio: g = g.replace(p, "")
                externo = "[" in g or "!" in g or "GETPIVOTDATA" in g.upper()
                cells[c.coordinate] = [None if externo else g, v, c.number_format]
            else:
                cells[c.coordinate] = [None, v if v is not None else f, c.number_format]
    larg = {k: d.width for k, d in wf.column_dimensions.items() if d.width}
    json.dump(dict(origem=f"{arq.split('/')[-1]} · aba {aba}", aba=aba, tipo=tipo, cells=cells, larg=larg),
              open(saida, "w"), ensure_ascii=False)
    print(len(cells), "células ·", sum(1 for x in cells.values() if x[0]), "fórmulas mantidas")


# ------------------------------------------------------------------ montagem na ferramenta
def _crono(wb):
    """Linhas-resumo (serviço '-') do cronograma: {(pacote, lote): (início, término)}."""
    ca = wb["CRONOGRAMA ATIVIDADES"]; out = {}
    for r in range(5, 6005):
        if ca.cell(r, 4).value == "-" and ca.cell(r, 3).value:
            out[(str(ca.cell(r, 3).value).strip(), str(ca.cell(r, 5).value or "").strip())] = (ca.cell(r, 6).value, ca.cell(r, 7).value)
    return out


def _resumo(col, pac, lote):
    """Data (F=início, G=término) da linha-resumo do pacote/lote no cronograma; 0 se não houver."""
    return (f'SUMPRODUCT(({CA}$C$5:$C$6004={pac})*({CA}$E$5:$E$6004={lote})*({CA}$D$5:$D$6004="-")*'
            f'{CA}${col}$5:${col}$6004)')


def _base(ws, A):
    for k, (f, v, nf) in A["cells"].items():
        c = ws[k]
        c.value = f if f else v
        if isinstance(c.value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}T00:00:00", c.value):
            c.value = dt.datetime.fromisoformat(c.value)
        if nf and nf != "General": c.number_format = nf
        c.font = Font(name=F_, size=9, color="000000" if f else "3C3A33")
    for k, w in A.get("larg", {}).items():
        ws.column_dimensions[k].width = w


def _h(ws, ref, txt, w=None):
    c = ws[ref]; c.value = txt; c.font = Font(name=F_, bold=True, size=9, color=NAVY); c.fill = fill("E8F0EA")
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    if w: ws.column_dimensions[re.match("[A-Z]+", ref)[0]].width = w


def montar(wb, A, R0, NM, cfg=None):
    """Cria a aba AUX.ESTRUTURA logo depois da ESTRUTURA. Devolve a função H(k) → fórmula da PREVISÃO INFORMADA do
    k-ésimo mês do horizonte (linha R0 + k da aba ESTRUTURA)."""
    cfg = cfg or {}
    pos = wb.sheetnames.index("ESTRUTURA") + 1 if "ESTRUTURA" in wb.sheetnames else len(wb.sheetnames)
    if NOME in wb.sheetnames: del wb[NOME]
    ws = wb.create_sheet(NOME, pos); ws.sheet_properties.tabColor = "C0A062"
    _base(ws, A)
    cr = _crono(wb)
    MES = [f"ESTRUTURA!$B${R0 + k}" for k in range(NM)]
    if A["tipo"] == "trechos":
        return _trechos(ws, A, cr, R0, NM, MES, cfg)
    return _pavimentos(ws, A, cr, R0, NM, MES, cfg)


def _trechos(ws, A, cr, R0, NM, MES, cfg):
    cells = A["cells"]
    pacotes = sorted({p for p, _ in cr}, key=len, reverse=True)
    rows = [int(re.sub(r"\D", "", k)) for k in cells if re.fullmatch(r"AC\d+", k)
            and cells[k][1] not in (None, "x") and int(re.sub(r"\D", "", k)) >= 3]
    r0, r1 = min(rows), max(rows)
    # pivô da origem (AP:AU) sai: a soma por mês agora é calculada abaixo
    for k in list(cells):
        col = re.match("[A-Z]+", k)[0]
        if CI(col) >= CI("AP") and CI(col) <= CI("AU"): ws[k].value = None
    _h(ws, "AW2", "LOTE NO CRONOGRAMA", 22); _h(ws, "AX2", "PACOTE NO CRONOGRAMA", 30)
    _h(ws, "AY2", "TÉRMINO NO CRONOGRAMA", 12); _h(ws, "AZ2", "TÉRMINO USADO", 12)
    _h(ws, "BA2", "TÉRMINO SIMULADO\n(digite)", 12); ws.column_dimensions["BB"].width = 3
    sim = {int(k): v for k, v in (cfg.get("simulacao") or {}).items()}
    ws.row_dimensions[2].height = 42
    sem = []
    for r in rows:
        v = str(cells[f"AC{r}"][1]).strip()
        pac = next((p for p in pacotes if v == p or v.endswith(" - " + p) or v.endswith("- " + p)), None)
        lote = v[:len(v) - len(pac)].rstrip(" -").strip() if pac else ""
        if pac and not lote:   # vínculo só com o pacote: lote único do pacote no cronograma
            ls = [l for p, l in cr if p == pac]; lote = ls[0] if len(ls) == 1 else ""
        if not pac: sem.append(v)
        ws[f"AW{r}"] = lote; ws[f"AX{r}"] = pac or ""
        ws[f"AY{r}"] = f'=IF(AX{r}="",0,{_resumo("G", f"AX{r}", f"AW{r}")})'
        ws[f"AZ{r}"] = f'=IF(ISNUMBER(BA{r}),BA{r},IF(AY{r}>0,AY{r},IF(ISNUMBER(AE{r}),AE{r},0)))'
        bc = ws[f"BA{r}"]; bc.fill = fill("FFF2CC"); bc.number_format = "dd/mm/yyyy"; bc.font = Font(name=F_, size=9, bold=True, color="0000FF")
        if r in sim:
            bc.value = dt.datetime.fromisoformat(sim[r]["fim"])
            bc.comment = Comment(f"Simulação: {sim[r].get('desc', '')}", "Ferramenta")
        for c in ("AY", "AZ"): ws[f"{c}{r}"].number_format = "dd/mm/yyyy"
        for c in ("AW", "AX", "AY", "AZ"): ws[f"{c}{r}"].font = Font(name=F_, size=9, color="1F7A52" if c in ("AY", "AZ") else "0000FF")
    ws["AW1"] = ("Datas: término simulado (coluna BA, se preenchido); senão o término da linha-resumo do pacote/lote na aba "
                 "CRONOGRAMA ATIVIDADES; sem par, vale a data da origem (coluna AE).")
    ws["AW1"].font = Font(name=F_, italic=True, size=9, color=INK2)
    # parâmetros de pagamento (aba ESTRUTURA da origem)
    P = cfg.get("pagamento", {"serv_forma": 0.7, "serv_aco": 1, "rt": 0.07, "aco_parc": 0.88, "corte_dobra": 0.12, "defasagem": 1})
    par = [("% serviço na forma (retenção)", "serv_forma"), ("% serviço no aço MDO (retenção)", "serv_aco"),
           ("Retenção técnica", "rt"), ("Aço material em 2 parcelas", "aco_parc"), ("Corte e dobra", "corte_dobra"),
           ("Pagamento: meses após o término", "defasagem")]
    _h(ws, "BQ2", "PARÂMETROS DE PAGAMENTO", 30); ws.column_dimensions["BR"].width = 9
    pr = {}
    for i, (t, k) in enumerate(par):
        ws[f"BQ{3 + i}"] = t; ws[f"BR{3 + i}"] = P[k]; pr[k] = f"$BR${3 + i}"
        ws[f"BQ{3 + i}"].font = Font(name=F_, size=9); ws[f"BR{3 + i}"].font = Font(name=F_, size=9, color="0000FF")
        ws[f"BR{3 + i}"].fill = fill("FFF2CC")
    # custo e pagamento por mês (horizonte da ferramenta)
    heads = [("BC", "MÊS", 9), ("BD", "FORMA", 11), ("BE", "RT FORMA", 10), ("BF", "AÇO MDO", 11), ("BG", "RT AÇO MDO", 10),
             ("BH", "SARRAF./ POLIMENTO", 11), ("BI", "CONCRETO", 11), ("BJ", "AÇO MAT", 11), ("BK", "AÇO 1ª PARCELA", 11),
             ("BL", "AÇO 2ª PARCELA", 11), ("BM", "CORTE E DOBRA", 11), ("BN", "CUSTO DO MÊS", 12), ("BO", "PAGAMENTO\n(curva da ESTRUTURA)", 13)]
    for c, t, w in heads: _h(ws, f"{c}2", t, w)
    rng = lambda c: f"${c}${r0}:${c}${r1}"
    dentro = lambda r: f'({rng("AZ")}>=BC{r})*({rng("AZ")}<=EOMONTH(BC{r},0))'
    for k in range(NM):
        r = 3 + k
        ws[f"BC{r}"] = f"={MES[k]}"; ws[f"BC{r}"].number_format = "mmm/yy"
        for c, src in (("BD", "AH"), ("BF", "AI"), ("BH", "AJ"), ("BI", "AL"), ("BJ", "AK")):
            ws[f"{c}{r}"] = f"=SUMPRODUCT({dentro(r)}*{rng(src)})"
        ws[f"BE{r}"] = f"=BD{r}*{pr['serv_forma']}*{pr['rt']}"
        ws[f"BG{r}"] = f"=BF{r}*{pr['serv_aco']}*{pr['rt']}"
        ws[f"BK{r}"] = f"=BJ{r}*{pr['aco_parc']}/2"
        ws[f"BL{r}"] = f"=BJ{r - 1}*{pr['aco_parc']}/2" if k else "=0"
        ws[f"BM{r}"] = f"=BJ{r}*{pr['corte_dobra']}"
        ws[f"BN{r}"] = f"=BD{r}-BE{r}+BF{r}-BG{r}+BH{r}+BI{r}+BK{r}+BL{r}+BM{r}"
        ws[f"BO{r}"] = f"=IF({k + 1}-{pr['defasagem']}<1,0,IFERROR(INDEX($BN$3:$BN${2 + NM},{k + 1}-{pr['defasagem']}),0))"
        for c in "BD BE BF BG BH BI BJ BK BL BM BN BO".split():
            ws[f"{c}{r}"].number_format = "#,##0;-#,##0;-"; ws[f"{c}{r}"].font = Font(name=F_, size=9, bold=c == "BO")
        ws[f"BC{r}"].font = Font(name=F_, size=9)
    rt = 3 + NM
    ws[f"BC{rt}"] = "TOTAL"; ws[f"BC{rt}"].font = Font(name=F_, bold=True, size=9)
    for c in "BD BE BF BG BH BI BJ BK BL BM BN BO".split():
        ws[f"{c}{rt}"] = f"=SUM({c}3:{c}{rt - 1})"; ws[f"{c}{rt}"].number_format = "#,##0"
        ws[f"{c}{rt}"].font = Font(name=F_, bold=True, size=9)
    ws["BC1"] = "Custo por mês de término dos trechos (regras de pagamento da aba ESTRUTURA da origem) → coluna BO = curva da ESTRUTURA"
    ws["BC1"].font = Font(name=F_, bold=True, size=10, color=NAVY)
    ws["BO2"].comment = Comment("A aba ESTRUTURA usa esta coluna como PREVISÃO INFORMADA (peso da distribuição do saldo "
                                "nos meses futuros). Pagamento = custo do mês deslocado pela defasagem (BR8).", "Ferramenta")
    ws.freeze_panes = "A3"
    return (lambda k: f"='{NOME}'!BO{3 + k}"), dict(trechos=len(rows), sem_par=sem, simulados=sorted(sim))


def _pavimentos(ws, A, cr, R0, NM, MES, cfg):
    cells = A["cells"]
    nome = {r: str(cells[f"H{r}"][1]).strip() for r in range(4, 41) if f"H{r}" in cells and cells[f"H{r}"][1]}
    # pavimento da origem → lote do pacote ESTRUTURA no cronograma
    pac = cfg.get("pacote", "ESTRUTURA")
    lotes = [l for p, l in cr if p == pac]
    def par(n):
        u = n.upper()
        m = re.match(r"(\d+)º", n)
        if m: return next((l for l in lotes if re.match(rf"{m.group(1)}º\s*PAV", l.upper())), "")
        for chave, alvo in (("COBERTURA", "COBERTURA GERAL"), ("BARRILETE", "BARRILETE"), ("CAIXA", "RESERVAT"),
                            ("INTERMEDI", "INTERMEDI"), ("TÉRREO", "TÉRREO"), ("SUBSOLO", "SUBSOLO")):
            if chave in u: return next((l for l in lotes if alvo in l.upper()), "")
        return ""
    # datas da origem: cabeçalhos AC2:AX2
    _h(ws, "AZ2", "LOTE NO CRONOGRAMA", 26); _h(ws, "BA2", "INÍCIO", 10); _h(ws, "BB2", "TÉRMINO", 10)
    ws.row_dimensions[2].height = 42
    C0 = CI("BD")
    for k in range(NM):
        c = L(C0 + k); ws[f"{c}3"] = f"={MES[k]}"; ws[f"{c}3"].number_format = "mmm/yy"
        ws[f"{c}3"].font = Font(name=F_, bold=True, size=9, color=NAVY); ws[f"{c}3"].fill = fill("E8F0EA")
        ws.column_dimensions[c].width = 10
    ws[f"{L(C0)}2"] = "CUSTO DE CADA PAVIMENTO POR MÊS (rateado pelos dias de execução no cronograma; sem par, distribuição da origem)"
    ws[f"{L(C0)}2"].font = Font(name=F_, bold=True, size=10, color=NAVY)
    ult = L(C0 + NM - 1)
    sem = []
    for r, n in nome.items():
        lote = par(n)
        if not lote: sem.append(n)
        ws[f"AZ{r}"] = lote; ws[f"AZ{r}"].font = Font(name=F_, size=9, color="0000FF")
        qp = '"' + pac + '"'
        ws[f"BA{r}"] = f'=IF(AZ{r}="",0,{_resumo("F", qp, f"AZ{r}")})'
        ws[f"BB{r}"] = f'=IF(AZ{r}="",0,{_resumo("G", qp, f"AZ{r}")})'
        for c in ("BA", "BB"): ws[f"{c}{r}"].number_format = "dd/mm/yy"; ws[f"{c}{r}"].font = Font(name=F_, size=9, color="1F7A52")
        for k in range(NM):
            c = L(C0 + k); m = f"{c}$3"
            ws[f"{c}{r}"] = (f"=IF(AND($BA{r}>0,$BB{r}>=$BA{r}),$AA{r}*MAX(0,MIN($BB{r},EOMONTH({m},0))-MAX($BA{r},{m})+1)/($BB{r}-$BA{r}+1),"
                             f"SUMPRODUCT((YEAR($AC$2:$AX$2)=YEAR({m}))*(MONTH($AC$2:$AX$2)=MONTH({m}))*$AC{r}:$AX{r}))")
            ws[f"{c}{r}"].number_format = "#,##0;-#,##0;-"; ws[f"{c}{r}"].font = Font(name=F_, size=9)
    rs = max(nome) + 1
    ws[f"AZ{rs}"] = "CURVA DA ESTRUTURA (soma)"; ws[f"AZ{rs}"].font = Font(name=F_, bold=True, size=9, color=NAVY)
    for k in range(NM):
        c = L(C0 + k)
        ws[f"{c}{rs}"] = f"=SUM({c}4:{c}{rs - 1})"; ws[f"{c}{rs}"].number_format = "#,##0;-#,##0;-"
        ws[f"{c}{rs}"].font = Font(name=F_, bold=True, size=9, color=NAVY); ws[f"{c}{rs}"].fill = fill("D9E6DC")
    ws[f"AZ{rs}"].comment = Comment("A aba ESTRUTURA usa esta linha como PREVISÃO INFORMADA (peso da distribuição do saldo "
                                    "nos meses futuros).", "Ferramenta")
    ws["AZ1"] = ("Datas: início e término da linha-resumo do pacote ESTRUTURA no lote do pavimento (aba CRONOGRAMA ATIVIDADES). "
                 "Lote vazio ou sem data: vale a distribuição manual da origem (colunas AC:AX).")
    ws["AZ1"].font = Font(name=F_, italic=True, size=9, color=INK2)
    ws.freeze_panes = "A4"
    return (lambda k: f"='{NOME}'!{L(C0 + k)}{rs}"), dict(pavimentos=len(nome), sem_par=sem)


if __name__ == "__main__" and sys.argv[1] == "extrair":
    extrair(*sys.argv[2:6])
