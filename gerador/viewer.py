import json, datetime as dt, re, html, sys
TPL, OUTF = (sys.argv[1], sys.argv[2]) if len(sys.argv) > 2 else ('viewer_tpl.html', '/home/user/silver-octo-meme/visualizar_ferramenta.html')
TOK = {'0000FF':'var(--in)','008000':'var(--link)','C00000':'var(--neg)','FFFFFF':'var(--hdrfg)','595959':'var(--mut)','FFFF00':'var(--yel)','D9E1F2':'var(--tot)','FFF2CC':'var(--inbg)','FCE4D6':'var(--red)','E2EFDA':'var(--ok)','1F4E78':'var(--hdr)','C65911':'var(--org)','FFC000':'var(--yel)'}
tk = lambda h: TOK.get(h.upper(), '#' + h)
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L
import os
V = load_workbook(os.environ.get('VW_VALUES', 'm.xlsx'), data_only=True)          # valores recalculados
Fw = load_workbook(os.environ.get('VW_FORMULAS', '/home/user/silver-octo-meme/fluxo_caixa_botanico_modelo.xlsx'))  # fórmulas e formatos
JSON_MODE = TPL == '--json'
MES = ['jan','fev','mar','abr','mai','jun','jul','ago','set','out','nov','dez']
def br(x, dec):
    s = f"{abs(x):,.{dec}f}".replace(',', 'X').replace('.', ',').replace('X', '.')
    return s
def fmt(v, nf):
    if v is None or v == '': return '', ''
    if isinstance(v, (dt.datetime, dt.date)):
        if 'dd' in nf: return v.strftime('%d/%m/%Y'), 'n'
        return f"{MES[v.month-1]}/{str(v.year)[2:]}" if 'mmm' in nf else v.strftime('%m/%Y'), 'n'
    if isinstance(v, bool): return str(v), ''
    if isinstance(v, (int, float)):
        if '%' in nf:
            dec = 2 if '0.00%' in nf else 1 if '0.0%' in nf else 0
            if abs(v) < 1e-9 and '"-"' in nf: return '-', 'n'
            return ('-' if v < 0 else '') + br(v*100, dec) + '%', 'n neg' if v < 0 else 'n'
        if nf in ('0', 'General') and float(v).is_integer() and 1900 <= v <= 2100: return str(int(v)), 'n'
        dec = 2 if '.00' in nf or nf == 'General' and not float(v).is_integer() else (1 if '0.0' in nf else (3 if '0.000' in nf else 0))
        if 'General' == nf and not float(v).is_integer(): dec = 2
        if abs(v) < 0.005 and '"-"' in nf: return '-', 'n'
        s = br(v, dec)
        if v < 0: return (f"({s})" if '(' in nf else '-' + s), 'n neg'
        return s, 'n'
    return str(v), ''
def color(c):
    try:
        if c is not None and c.type == 'rgb' and c.rgb and isinstance(c.rgb, str): return c.rgb[-6:]
    except Exception: pass
    return None
styles = {}
def sid(fc):
    f = fc.font; fill = fc.fill
    key = (color(f.color) if f else None, bool(f and f.b), bool(f and f.i),
           color(fill.fgColor) if fill and fill.fill_type == 'solid' else None)
    if key not in styles: styles[key] = len(styles)
    return styles[key]
SKIP = {'Tabela', 'PREMISSAS'}
LIMIT = {'CURVA PREVISION': (60, 30), 'DADOS BRUTOS': (60, 90), 'CRONOGRAMA ATIVIDADES': (120, 27)}
out = []
for ws in Fw.worksheets:
    if ws.title in SKIP: continue
    vs = V[ws.title]
    hidden = {k for k, d in ws.column_dimensions.items() if d.hidden}
    maxr, maxc = ws.max_row, ws.max_column
    if ws.title in LIMIT: maxr, maxc = min(maxr, LIMIT[ws.title][0]), min(maxc, LIMIT[ws.title][1])
    if ws.title == 'DADOS HISTORICOS': maxc = 77
    cols = [c for c in range(1, maxc+1) if L(c) not in hidden]
    widths = [round((ws.column_dimensions[L(c)].width or 9) * 8) for c in cols]
    rows = []
    for r in range(1, maxr+1):
        cells = []
        for j, c in enumerate(cols):
            fc = ws.cell(r, c); v = vs.cell(r, c).value
            txt, cls = fmt(v, fc.number_format or 'General')
            f = fc.value if isinstance(fc.value, str) and fc.value.startswith('=') else ''
            s = sid(fc)
            if txt == '' and s == 0 and not f: continue
            cells.append([j, txt, cls, s, f])
        rows.append(cells)
    while rows and not rows[-1]: rows.pop()
    out.append({'name': ws.title, 'cols': [L(c) for c in cols], 'w': widths, 'rows': rows,
                'freeze': ws.freeze_panes, 'note': 'Mostrando só as primeiras linhas/colunas.' if ws.title in LIMIT else ''})
css = []
for (fc, b, i, bg), k in styles.items():
    rule = []
    if fc and fc != '000000': rule.append(f'color:{tk(fc)}')
    elif bg: rule.append('color:var(--oncolor)')
    if b: rule.append('font-weight:700')
    if i: rule.append('font-style:italic')
    if bg: rule.append(f'background:{tk(bg)}')
    css.append(f'.s{k}{{{";".join(rule)}}}')
data = json.dumps(out, ensure_ascii=False, separators=(',', ':'))
if not JSON_MODE:
    tpl = open(TPL).read()
    open(OUTF, 'w').write(tpl.replace('/*CSS*/', '\n'.join(css)).replace('/*DATA*/', data))
print(len(data)//1024, 'KB', len(out), 'abas')

# ---------------- dados do painel (DASHBOARD) para a visualização online
def dash_data():
    d = V['DASHBOARD']; vp = V['VALOR POR PL']; bi = V['MODELO PARA BI']
    def num(x): return float(x) if isinstance(x, (int, float)) else 0.0
    years = []
    for r in range(13, 20):
        y = d.cell(r, 2).value
        years.append(dict(y=int(y), anual=num(d.cell(r, 3).value), acum=num(d.cell(r, 4).value), pa=num(d.cell(r, 5).value),
                          pac=num(d.cell(r, 6).value), real=num(d.cell(r, 8).value), proj=num(d.cell(r, 9).value)))
    groups = [dict(name=d.cell(r, 2).value, v=num(d.cell(r, 4).value), p=num(d.cell(r, 5).value),
                   av=(None if d.cell(r, 6).value in (None, '') else num(d.cell(r, 6).value))) for r in range(26, 32)]
    curve = [dict(m=bi.cell(1, c).value.strftime('%Y-%m'), p=num(bi.cell(65, c).value), v=num(bi.cell(63, c).value)) for c in range(8, 80)]
    pend = [f"{vp.cell(r, 4).value}: {vp.cell(r, 14).value.lstrip('⚠ ')}" for r in range(11, 49) if vp.cell(r, 14).value not in (None, 'OK')]
    pend += [f"{vp.cell(r, 4).value}: {vp.cell(r, 6).value.lstrip('⚠ ')}" for r in range(54, 60)
             if isinstance(vp.cell(r, 6).value, str) and vp.cell(r, 6).value.startswith('⚠') and r not in (54, 59)]
    corte = vp['P5'].value
    return dict(title=d['B2'].value, sub=d['B3'].value, month=d['B11'].value, ano=int(V['RETRATO 2027']['D5'].value),
                corte=corte.strftime('%Y-%m'), proj=num(d['B6'].value), real=num(d['D6'].value), preal=num(d['E6'].value),
                ano_v=num(d['F6'].value), saldo=num(d['H6'].value), conf=d['I6'].value, years=years,
                total=num(d.cell(20, 3).value), treal=num(d.cell(20, 8).value), tproj=num(d.cell(20, 9).value),
                groups=groups, gtotal=num(d.cell(32, 4).value), curve=curve, pend=pend)
def ret_data():
    r = V['RETRATO 2027']
    def num(x): return float(x) if isinstance(x, (int, float)) else None
    rows = []
    for rr in range(12, 80):
        nm = r.cell(rr, 3).value
        if nm is None: continue
        kind = 'tot' if isinstance(nm, str) and nm.startswith('DESEMBOLSO DO ANO') else ('sub' if r.cell(rr, 2).value and ' a ' in str(r.cell(rr, 2).value) else 'pl')
        rows.append(dict(k=kind, c=r.cell(rr, 2).value, n=nm, ff=num(r.cell(rr, 4).value), fc=num(r.cell(rr, 5).value), rs=num(r.cell(rr, 6).value),
                         ic=num(r.cell(rr, 7).value), rsel=num(r.cell(rr, 8).value), pa=num(r.cell(rr, 9).value), fin=num(r.cell(rr, 10).value),
                         obs=r.cell(rr, 11).value or '', nota=r.cell(rr, 12).value or ''))
    return dict(title=r['B2'].value, sub=r['B3'].value, ano=int(r['D5'].value), incc_idx=num(r['G5'].value),
                incc_mes=r['F5'].value.strftime('%m/%Y'), cards=[dict(l=r[f'{c}7'].value, v=r[f'{c}8'].value, n=r[f'{c}9'].value) for c in 'BDFHJK'], rows=rows)
if JSON_MODE:
    dd = dash_data(); rr = ret_data()
    obra = json.dumps(dict(css='\n'.join(css), D=out, DASH=dd, RET=rr), ensure_ascii=False, default=str, separators=(',', ':'))
    open(OUTF, 'w').write(obra)
    card = dict(nome=os.environ.get('VW_NOME', 'Obra'), id=os.environ.get('VW_ID', 'obra'), corte=dd['corte'], ano=dd['ano'], proj=dd['proj'],
                real=dd['real'], preal=dd['preal'], ano_v=dd['ano_v'], saldo=dd['saldo'], conf=dd['conf'],
                gerado=dt.date.today().isoformat(), kb=round(len(obra) / 1024))
    open(OUTF + '.card.json', 'w').write(json.dumps(card, ensure_ascii=False))
else:
    _h = open(OUTF).read().replace('/*DASH*/null', json.dumps(dash_data(), ensure_ascii=False)).replace('/*RET*/null', json.dumps(ret_data(), ensure_ascii=False, default=str))
    open(OUTF, 'w').write(_h)
