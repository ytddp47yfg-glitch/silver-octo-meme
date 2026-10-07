"""Dados do painel GANTT a partir da aba CRONOGRAMA ATIVIDADES (export do Prevision) da ferramenta.

Uso: python3 gantt.py arquivo_calc.xlsx saida.json

Árvore: PL (coluna M da aba, código do VALOR POR PL) → pacote de trabalho → lote/pavimento (linha-resumo do
Prevision, serviço '-') → tarefas (demais linhas do mesmo pacote e lote). Pais: início = menor início, término =
maior término, % = média dos filhos ponderada pela duração.
"""
import sys, json, re, unicodedata, datetime as dt
from openpyxl import load_workbook


# ---------------- avaliação: pacote descrito de forma expressa na curva do Prevision?
def norm(s): return unicodedata.normalize('NFKD',str(s).upper()).encode('ascii','ignore').decode()
STOP=set('DE DA DO DAS DOS E EM PARA COM A O AS OS NO NA NOS NAS POR PAV SETOR TRECHO FASE ETAPA PARTE TIPO GERAL INSTALACAO INSTALACOES INST MONTAGEM EXECUCAO SERVICO SERVICOS OBRA OBRAS'.split())
LOC=set('''TORRE EMBASAMENTO MEDICAO LOOKAHEAD LOOKAEHAD LOOKAREAD GARDEN POSTERIOR LATERAL ESQUERDA DIREITA FRONTAL AREA SOCIAL QUADRANTE INTERNO INTERNA EXTERNO EXTERNA
TETO PAREDES PAREDE FINAL FINA GROSSA BALTAZAR JAIRO CREMALHEIRA RAMPA CASA MAQUINAS TERRACO SALAO FESTA DEMAO BLOCO SUBSOLO COBERTURA TERREO PILOTIS LAZER BARRILETE
VARANDA VARANDAS APTO APTOS DUPLEX PROJECAO ALAMEDA CONCORDIA OLYMPUS OLIMPUS ACACIA CHOPP FABRICA PLANO ACAO PREDIO ULTIMA PRIMEIRA PRACA TRABALHO FACE RUA BOM JESUS
IPE MANACA IGREJA FUNDO FUNDOS LIBERACAO POS INTER SEGUNDA TERCEIRA'''.split())
SIN={'HIDRAULICA':'HIDRO','HIDRAULICAS':'HIDRO','HIDRAULICO':'HIDRO','HIDROSSANITARIA':'HIDRO','HIDROSSANITARIAS':'HIDRO','HIDROSANITARIA':'HIDRO','HIDROSANITARIAS':'HIDRO','HIDRAUCAS':'HIDRO','HIDRAULICAS.':'HIDRO'}
GEN=set('INFRA INFRAESTRUTURA'.split())   # genéricos: não contam como serviço principal
def toks(s): return [SIN.get(t,t) for t in re.findall(r'[A-Z]{2,}',norm(s)) if t not in STOP and len(t)>=3 or t in ('GAS',)]
def st(t): return t[:5]
def avaliar(wb):
    """Pacote do cronograma de atividades × curva do Prevision da PL associada (EAP da aba CURVA PREVISION).
    ok=True quando o serviço principal do pacote aparece na EAP da PL; senão devolve a nota explicando a associação."""
    cv=wb['CURVA PREVISION']; ca=wb['CRONOGRAMA ATIVIDADES']; cps=wb['CP RESUMO']; vp=wb['VALOR POR PL']
    lin={}
    for r in range(5,60):
        a,l=cps.cell(r,1).value,cps.cell(r,3).value
        if a and isinstance(l,(int,float)): lin[str(a).strip()]=(str(cv.cell(int(l),2).value),str(cv.cell(int(l),3).value))
    cods=[(str(cv.cell(r,2).value),str(cv.cell(r,3).value)) for r in range(4,cv.max_row+1) if cv.cell(r,2).value and cv.cell(r,3).value]
    q={str(vp.cell(r,3).value).strip():str(vp.cell(r,17).value).strip() for r in range(11,49) if vp.cell(r,3).value and vp.cell(r,17).value}
    desc=dict(cods)
    sub={};nome={}
    for pl,pc in q.items():
        # EAP própria da PL (código do Prevision na coluna Q), mesmo que a curva siga outro item (coluna R / CP RESUMO)
        if pc in desc: e,nm=pc,desc[pc]
        elif pc in lin: e,nm=lin[pc]
        else: continue
        nome[pl]=(e,nm)
        sub[pl]=[(c,d) for c,d in cods if c==e or c.startswith(e+'.')]
    stems={pl:set(st(t) for c,d in v for t in toks(d)) for pl,v in sub.items()}
    pk={}
    for r in range(5,6005):
        p=ca.cell(r,3).value
        if p: pk.setdefault(str(p).strip(),str(ca.cell(r,13).value or '').strip())
    out={}
    for p,pl in pk.items():
        core=[t for t in toks(p) if t not in LOC and t not in GEN] or [t for t in toks(p) if t not in LOC] or toks(p)
        if not pl or pl not in sub:
            out[p]=dict(ok=False,nota='Pacote sem PL correspondente no orçamento/Prevision (atividade de gestão ou entrega): não entra em nenhuma curva de PL.'); continue
        e,nm=nome[pl]
        achou=[t for t in core if st(t) in stems[pl]]
        if core and st(core[0]) in stems[pl]:
            item=next((f"{c} {d}" for c,d in sub[pl] if st(core[0]) in {st(x) for x in toks(d)}),'')
            out[p]=dict(ok=True,item=item)
        else:
            falta=[t for t in core if st(t) not in stems[pl]]
            outras=sorted({k for k,v in stems.items() if k!=pl and core and st(core[0]) in v})
            nt=(f"Serviço não descrito de forma expressa na curva do Prevision da PL {e} – {nm}"
                f" (termo sem correspondência: {', '.join(t.lower() for t in falta[:3])}). Associado a esta PL por afinidade do serviço (regra de de-para da ferramenta).")
            if outras: nt+=" O termo aparece na curva de: "+"; ".join(f"{nome[k][0]} – {nome[k][1]}" for k in outras[:3])+"."
            out[p]=dict(ok=False,nota=nt)
    return out


def main():
    arq, out = sys.argv[1], sys.argv[2]
    wb = load_workbook(arq, data_only=True)
    ws = wb['CRONOGRAMA ATIVIDADES']; vp = wb['VALOR POR PL']
    nome_pl = {str(vp.cell(r, 3).value).strip(): vp.cell(r, 4).value for r in range(11, 49) if vp.cell(r, 3).value}
    ref = None
    rows = []
    for r in range(5, 6005):
        pac = ws.cell(r, 3).value
        if not pac: continue
        ini, fim = ws.cell(r, 6).value, ws.cell(r, 7).value
        if not isinstance(ini, dt.datetime) or not isinstance(fim, dt.datetime): continue
        ref = ref or ws.cell(r, 8).value
        num = lambda c: float(ws.cell(r, c).value) if isinstance(ws.cell(r, c).value, (int, float)) else 0.0
        rows.append(dict(pac=str(pac).strip(), srv=str(ws.cell(r, 4).value or '-').strip(), lote=str(ws.cell(r, 5).value or '').strip(),
                         i=ini.date(), f=fim.date(), pb=num(9), pp=num(10), pr=num(11), pl=str(ws.cell(r, 13).value or '').strip(),
                         pav=ws.cell(r, 14).value, sit=ws.cell(r, 15).value or ''))
    aval = avaliar(wb)
    ordem = {}
    arvore = {}
    for x in rows:   # PL → pacote → lote → tarefas, na ordem do cronograma
        pl = arvore.setdefault(x['pl'], {})
        pk = pl.setdefault(x['pac'], {})
        lt = pk.setdefault(x['lote'], {'resumo': None, 'tarefas': []})
        if x['srv'] == '-' and lt['resumo'] is None: lt['resumo'] = x
        else: lt['tarefas'].append(x)
        ordem.setdefault(x['pl'], len(ordem))

    def folha(x, nome):
        return dict(n=nome, i=x['i'].isoformat(), f=x['f'].isoformat(), pp=round(x['pp'], 1), pr=round(x['pr'], 1),
                    pb=round(x['pb'], 1), s=x['sit'])

    def junta(nome, filhos, extra=None):
        i = min(c['i'] for c in filhos); f = max(c['f'] for c in filhos)
        dur = lambda c: max(1, (dt.date.fromisoformat(c['f']) - dt.date.fromisoformat(c['i'])).days)
        sw = sum(dur(c) for c in filhos)
        no = dict(n=nome, i=i, f=f, pp=round(sum(c['pp'] * dur(c) for c in filhos) / sw, 1),
                  pr=round(sum(c['pr'] * dur(c) for c in filhos) / sw, 1), pb=round(sum(c['pb'] * dur(c) for c in filhos) / sw, 1), c=filhos)
        if extra: no.update(extra)
        return no

    raiz = []
    for pl, pacs in sorted(arvore.items(), key=lambda kv: min(x['i'] for p in kv[1].values() for l in p.values()
                                                               for x in ([l['resumo']] if l['resumo'] else []) + l['tarefas'])):
        nos_pac = []
        for pac, lotes in pacs.items():
            nos_lote = []
            for lote, d in lotes.items():
                tarefas = [folha(t, t['srv']) for t in d['tarefas']]
                if d['resumo']:
                    no = folha(d['resumo'], lote or pac)
                    if tarefas: no['c'] = tarefas
                elif tarefas:
                    no = junta(lote or pac, tarefas)
                else: continue
                nos_lote.append(no)
            nos_lote.sort(key=lambda c: c['i'])
            av = aval.get(pac, {'ok': True})
            nos_pac.append(nos_lote[0] if len(nos_lote) == 1 and not nos_lote[0].get('c') and nos_lote[0]['n'] in ('GERAL', pac)
                           else junta(pac, nos_lote))
            if nos_pac[-1]['n'] != pac: nos_pac[-1] = dict(nos_pac[-1], n=pac)
            if not av['ok']: nos_pac[-1] = dict(nos_pac[-1], x=1, nt=av['nota'])
        nos_pac.sort(key=lambda c: c['i'])
        nx = sum(1 for c in nos_pac if c.get('x'))
        raiz.append(junta(nome_pl.get(pl, pl or 'SEM PL (gestão / entrega)'), nos_pac, {'cod': pl, **({'nx': nx} if nx else {})}))
    corte = vp['P5'].value
    dados = dict(ref=(ref.date().isoformat() if isinstance(ref, dt.datetime) else None), corte=corte.strftime('%Y-%m'),
                 ret=f"{wb['RETRATO 2027']['D5'].value}-12-31", ini=min(x['i'] for x in rows).isoformat(),
                 fim=max(x['f'] for x in rows).isoformat(), n=len(rows), arv=raiz)
    json.dump(dados, open(out, 'w'), ensure_ascii=False, separators=(',', ':'))
    print(sum(1 for v in aval.values() if not v['ok']), 'pacotes sem descrição expressa na curva ·', len(rows), 'atividades ·', len(raiz), 'PLs ·', sum(len(p['c']) for p in raiz), 'pacotes · referência', dados['ref'], '·', dados['ini'], '→', dados['fim'])


if __name__ == '__main__':
    main()
