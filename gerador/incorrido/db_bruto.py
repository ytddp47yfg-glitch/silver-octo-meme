"""Monta DADOS BRUTOS (grupos de CUSTOS COM OBRAS × meses, custos positivos) a partir da aba DADOS (lançamentos).
Uso: db_bruto.py arquivo.xlsx 'GRUPO 1[;...]' '[VALOR REAL]|[VLR MOVIMENTO]' data2_atual.json saida.json
Exclui as projeções carregadas no ERP (A PAGAR sem origem no extrato)."""
import sys; sys.path.insert(0, '/tmp/claude-0/-home-user-silver-octo-meme/252de686-5b3c-58a7-a16e-5b6ca412b5a3/scratchpad/pylib')
import openpyxl, collections, json, re, datetime as dt
from openpyxl.utils import get_column_letter as L
arq, g1, val, d2p, out = sys.argv[1:6]
ws = openpyxl.load_workbook(arq, data_only=True, read_only=True)['DADOS']
rows = list(ws.iter_rows(values_only=True)); h = rows[0]; ix = {x: i for i, x in enumerate(h) if x}
agg = collections.defaultdict(lambda: collections.defaultdict(float)); excl = 0.0
for r in rows[1:]:
    if not r or r[ix['[GRUPO 1]']] not in g1.split(';') or r[ix['[PL GRUPO G3]']] != 'CUSTOS COM OBRAS': continue
    m, v = r[ix['[INÍCIO DO MÊS]']], r[ix[val]]
    if not isinstance(m, dt.datetime) or not isinstance(v, (int, float)): continue
    if r[ix['[STATUS FINANCEIRO]']] == 'A PAGAR' and r[ix['[ORIGEM EXTRATO]']] is None: excl += -v; continue
    agg[r[ix['[PL GRUPO G1]']]][m.date()] += -v
old = json.load(open(d2p))['DADOS BRUTOS']
ordem = [v for k, v in sorted(((int(k[1:]), v) for k, v in old.items() if re.fullmatch(r'A\d+', k)))]
grupos = [g for g in ordem if g in agg] + sorted(g for g in agg if g not in ordem)
meses = sorted({m for g in agg.values() for m, v in g.items() if abs(v) > 0.001})
DB = {'A1': f'Soma de {val} (sinal invertido: custos positivos) · aba DADOS, {g1}, CUSTOS COM OBRAS, sem projeções (A PAGAR sem origem)',
      'B1': 'Rótulos de Coluna', 'A2': 'Rótulos de Linha', 'A3': 'CUSTOS COM OBRAS', 'A4': 'MATERIAIS E SERVIÇOS PRESTADOS'}
for k, m in enumerate(meses): DB[f'{L(2 + k)}2'] = float((m - dt.date(1899, 12, 30)).days)
for i, g in enumerate(grupos):
    DB[f'A{5 + i}'] = g
    for k, m in enumerate(meses):
        v = agg[g].get(m, 0)
        if abs(v) > 0.001: DB[f'{L(2 + k)}{5 + i}'] = round(v, 2)
json.dump(DB, open(out, 'w'), ensure_ascii=False)
tot = lambda f: sum(v for g in agg.values() for m, v in g.items() if f(m))
print(f'{len(grupos)} grupos, {len(meses)} meses ({meses[0]:%m/%Y} a {meses[-1]:%m/%Y}) | até ago/26 R$ {tot(lambda m: m <= dt.date(2026, 8, 1)):,.2f} | set/26 R$ {tot(lambda m: m == dt.date(2026, 9, 1)):,.2f} | depois R$ {tot(lambda m: m > dt.date(2026, 9, 1)):,.2f} | projeções excluídas R$ {excl:,.2f}')
