"""Lê a tabela dinâmica do ERP (grupos × meses), pega os grupos de CUSTOS COM OBRAS, inverte o sinal e
compara com a aba DADOS BRUTOS atual da ferramenta. Uso: novo_db.py arquivo.xlsx aba data2.json saida.json"""
import sys; sys.path.insert(0, '/tmp/claude-0/-home-user-silver-octo-meme/252de686-5b3c-58a7-a16e-5b6ca412b5a3/scratchpad/pylib')
import openpyxl, json, re, datetime as dt, collections
from openpyxl.utils import get_column_letter as L
arq, aba, d2p, out = sys.argv[1:5]
ws = openpyxl.load_workbook(arq, data_only=True, read_only=True)[aba]
rows = list(ws.iter_rows(values_only=True))
hi = next(i for i, r in enumerate(rows) if r and r[0] == 'Rótulos de Linha'); hdr = rows[hi]
i0 = next(i for i in range(hi + 1, len(rows)) if rows[i][0] == 'CUSTOS COM OBRAS')
FIM = {'DEDUÇÕES DA RECEITA', 'DESPESAS EMPREENDIMENTO', 'DISTRIBUIÇÕES', 'EMPRÉSTIMOS', 'IMPOSTOS SOBRE FATURAMENTO', 'RECEITA BRUTA',
       'RESULTADO FINANCEIRO', 'TERRENO', 'TRANSFERÊNCIAS', 'Total Geral'}
i1 = next(i for i in range(i0 + 1, len(rows)) if rows[i][0] in FIM)
cr = rows[i0:i1]
mcols = [j for j in range(1, len(hdr)) if isinstance(hdr[j], dt.datetime)]
cols = [j for j in mcols if any(isinstance(r[j], (int, float)) and abs(r[j]) > 0.001 for r in cr[2:])]
DB = {'A1': f'{rows[hi - 1][0]} (sinal invertido: custos positivos) · {aba}', 'B1': 'Rótulos de Coluna', 'A2': 'Rótulos de Linha'}
for k, j in enumerate(cols): DB[f'{L(2 + k)}2'] = float((hdr[j] - dt.datetime(1899, 12, 30)).days)
new = collections.defaultdict(lambda: collections.defaultdict(float))
for i, r in enumerate(cr):
    DB[f'A{3 + i}'] = r[0]
    for k, j in enumerate(cols):
        v = r[j]
        if isinstance(v, (int, float)) and abs(v) > 0.001:
            DB[f'{L(2 + k)}{3 + i}'] = round(-v, 2)
            if i >= 2: new[r[0]][hdr[j].date()] += -v
json.dump(DB, open(out, 'w'), ensure_ascii=False)
# atual
old_db = json.load(open(d2p))['DADOS BRUTOS']
oh = {re.match(r'[A-Z]+', k)[0]: v for k, v in old_db.items() if re.fullmatch(r'[A-Z]+2', k) and isinstance(v, (int, float))}
old = collections.defaultdict(lambda: collections.defaultdict(float)); lab = {}
for k, v in old_db.items():
    m = re.fullmatch(r'([A-Z]+)(\d+)', k); c, r = m.group(1), int(m.group(2))
    if c == 'A': lab[r] = v
for k, v in old_db.items():
    m = re.fullmatch(r'([A-Z]+)(\d+)', k); c, r = m.group(1), int(m.group(2))
    if r >= 3 and c != 'A' and c in oh and isinstance(v, (int, float)) and lab.get(r) not in ('CUSTOS COM OBRAS', 'MATERIAIS E SERVIÇOS PRESTADOS', 'APORTES'):
        old[lab[r]][(dt.datetime(1899, 12, 30) + dt.timedelta(days=oh[c])).date()] += v
tot = lambda D: sum(sum(x.values()) for x in D.values())
ult = lambda D: max((m for g in D.values() for m, v in g.items() if abs(v) > 0.5 and m <= dt.date(2026, 10, 1)), default=None)
print(f'{len(cr) - 2} grupos, {len(cols)} meses ({hdr[cols[0]]:%m/%Y} a {hdr[cols[-1]]:%m/%Y}) | total novo {tot(new):,.2f} | atual {tot(old):,.2f} | dif {tot(new) - tot(old):,.2f}')
meses = sorted({m for D in (new, old) for g in D.values() for m in g})
print('por mês (só onde mudou, R$ mil):', [(f'{m:%m/%y}', round((sum(new[g].get(m, 0) for g in new) - sum(old[g].get(m, 0) for g in old)) / 1e3, 1)) for m in meses
       if abs(sum(new[g].get(m, 0) for g in new) - sum(old[g].get(m, 0) for g in old)) > 1])
nomes_novos = [g for g in new if g not in old and abs(sum(new[g].values())) > 0.5]; nomes_sumidos = [g for g in old if g not in new and abs(sum(old[g].values())) > 0.5]
print('grupos novos:', nomes_novos, '| grupos que sumiram:', nomes_sumidos)
print('por grupo (R$ mil, |dif|>1):', [(g[:30], round((sum(new[g].values()) - sum(old.get(g, {}).values())) / 1e3, 1)) for g in new if abs(sum(new[g].values()) - sum(old.get(g, {}).values())) > 1000])
