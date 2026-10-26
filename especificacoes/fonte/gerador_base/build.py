import json, re, copy, math
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

import os, tempfile
S = os.path.dirname(os.path.abspath(__file__))  # pasta do gerador base
TMP = tempfile.mkdtemp()  # imagens intermediárias
MODEL = os.path.join(S, '..', 'entrada', 'CDJ-GRE-ARQ-AC-R00_MODELO.xlsx')
# grava em saída separada para não sobrescrever a grade atual (que recebeu os scripts de especificacoes/gerador)
OUT = os.path.join(S, 'saida', 'CDJ-GRE-ARQ-AC-R00_base.xlsx'); os.makedirs(os.path.dirname(OUT), exist_ok=True)
FONTE = 'Tabela de acabamentos Ed. Cidade Jardim v00 (19/06/2026)'
DATA = '30/09/2026'
REV = 'REV. R00'

recs = json.load(open(f'{S}/recs.json'))
nm = lambda s: ' '.join(s.split())

# correção de texto sobreposto no PDF (célula D25)
for r in recs:
    r['PEDRAS'] = ['D25 - Bancada e cuba esculpida com bandeja removível em pedra mármore Barrocci Honed, '
                   'acabamento levigado, espessura 2cm. Ref.: Basigran' if v.startswith('D25') else v
                   for v in r['PEDRAS']]

# os dois vestiários aparecem com o mesmo nome na tabela (pág. 2); a pág. 8 (VESTIÁRIOS - TERMAS)
# corresponde ao de piso P10
vest = [r for r in recs if r['amb'].startswith('Vestiário')]
vest[0]['amb'] = 'VEST_TERMAS'   # piso P10
vest[1]['amb'] = 'VEST_PISCINA'  # piso P7
R = {r['amb']: r for r in recs}

# detalhe da pág. 7 (TERMAS) para o mosaico Bisazza
R['Thermas']['PAREDE'] = [
    'Painel: PA49 - Revestimento em Mosaico Bisazza 20x20cm, assentado com argamassa colante em pó ACIII. '
    'Fundo: NARCISO Shading Blend - 20x20mm. Desenho personalizado: NEW STERLIZIA Shading Blend - 10x10mm. '
    'Rejunte Fillgel Plus - 1101 - cor Bianco Alabastro (para o New Sterlizia a Bisazza recomenda Fillgel Plus - 6601 - cor rosa perla). '
    'Fabricante do rejunte: Drylevis',
    'Pilar: PA49 - Revestimento em Mosaico Bisazza 20x20cm. Fundo: NARCISO Shading Blend - 20x20mm. '
    'Rejunte Fillgel Plus - 1101 - cor Bianco Alabastro. Fabricante do rejunte: Drylevis',
    'PA46 - Revestimento de mármore branco selecionado, espessura 2cm, assentado com argamassa de cimento e areia. '
    'Rejunte em massa plástica a ser tonalizado de acordo com a cor da pedra',
    'PA48 - Revestimento de mármore branco selecionado, acabamento frisado com ranhuras de 4x4mm espassada a cada 12mm, '
    'assentado com argamassa de cimento e areia',
]

COL = dict(PISO='B', SOLEIRA='C', FILETE='D', PEITORIL='E', RODAPE='F', SOCULO='G', PAREDE='H', TETO='I',
           BANCADA='J', OUTROS_REV='K', LOUCAS='L', VALVULAS='M', SIFOES='N', TORNEIRAS='O', BACIAS='P',
           DUCHAS='Q', REGISTROS='R', ACION_CHUV='S', RALOS='T', OUTROS_LM='U', ELETRICOS='V', MACANETAS='W',
           PORTAS='X', JANELAS='Y', OUTROS_COMP='Z')

split_items = lambda s: [nm(x) for x in re.split(r'\|\s*(?=\*{0,3}(?:D\d+|\d+)\s*-)|\s+(?=D\d+\s+-\s)', s) if nm(x)]

def classify(rec):
    out = {}
    add = lambda k, v: out.setdefault(k, []).append(v)
    for v in rec['PISO']:
        if v != '-': add('PISO', v)
    for v in rec['RODAPÉ']:
        if v != '-': add('RODAPE', v)
    for v in rec['TETO']:
        if v != '-': add('TETO', v)
    for v in rec['PAREDE']:
        if v == '-': continue
        if v.startswith('P28 (D43)'): add('PORTAS', v)
        elif re.match(r'D\d+ ', v) or v.startswith('Perfil'): add('OUTROS_REV', v)
        else: add('PAREDE', v)
    for v in rec['PEDRAS']:
        if v == '-': continue
        l = v.lower()
        if 'bancada' in l: add('BANCADA', v)
        else: add('OUTROS_REV', v)
    for src in ('LOUÇAS E METAIS', 'DIVERSOS'):
        for cell in rec[src]:
            if cell == '-': continue
            for it in split_items(cell):
                l = it.lower()
                if 'bacia' in l: add('BACIAS', it)
                elif 'monocomando para chuveiro' in l: add('ACION_CHUV', it)
                elif 'chuveiro' in l: add('OUTROS_LM', it)
                elif 'torneira' in l or 'monocomando' in l or 'misturador' in l: add('TORNEIRAS', it)
                elif 'cuba' in l or 'tanque' in l: add('LOUCAS', it)
                elif 'ralo' in l: add('RALOS', it)
                else: add('OUTROS_COMP', it)
    return out

# ambiente da tabela -> linhas do modelo (aba PILOTIS)
MAP = [
    ('Lobby', ['LOBBY PILOTIS']),
    ('Salão de Festas', ['SALÃO DE FESTAS']),
    ('IS Fem| Masc - Salão de Festas', ['I.S FEMININO', 'I.S MASCULINO']),
    ('Cozinha salão de festas', ['COZINHA SALÃO DE FESTAS']),
    ('Dml', ['DML']),
    ('IS| IS PCD - Lobby', ['I.S P.C.D', 'I.S']),
    ('Espaço Kids', ['ESPAÇO KIDS']),
    ('Espaço Gourmet', ['ESPAÇO GOURMET']),
    ('Piscina descoberta', ['PISCINA DECOBERTA']),
    ('Deck piscina', ['DECK']),
    ('VEST_PISCINA', ['VESTIÁRIO FEMININO', 'VESTIÁRIO MASCULINO']),
    ('Academia', ['ACADEMIA']),
    ('Spa| Piscina Raia', ['PISCINA RAIA', 'SPA']),
    ('Thermas', ['ESPAÇO TERMAS']),
    ('Sauna a vapor', ['SAUNA VAPOR']),
    ('Sauna seca', ['SAUNA SECA']),
    ('Massagem', ['MASSAGEM']),
    ('Espaço da beleza', ['BELEZA']),
    ('Quadra de tênis', ['QUADRA TÊNIS']),
    ('Depósito| Depósito jardim', ['DEP JARDIM', 'DEPÓSITO']),
]
NEW = [  # ambientes da tabela sem linha no modelo
    ('VEST_TERMAS', 'VESTIÁRIO FEM./MASC. - TERMAS'),
    ('Escada de incêndio', 'ESCADA DE INCÊNDIO'),
    ('Antecâmara', 'ANTECÂMARA'),
    ('Circulação', 'CIRCULAÇÃO'),
    ('Estar Piscina| Circulação', 'ESTAR PISCINA / CIRCULAÇÃO'),
    ('Piso - Circ. Externa coberta| descoberta e Varanda Espaço kids', 'CIRCULAÇÃO EXTERNA COBERTA/DESCOBERTA E VARANDA ESPAÇO KIDS'),
    ('Escada externa', 'ESCADA EXTERNA'),
    ('Quadra Infantil', 'QUADRA INFANTIL'),
    ('Apoio quadras', 'APOIO QUADRAS'),
]
assert set(R) == {a for a, _ in MAP} | {a for a, _ in NEW}, set(R) ^ ({a for a, _ in MAP} | {a for a, _ in NEW})

wb = openpyxl.load_workbook(MODEL)
ws = wb['PILOTIS']
rows_by_name = {}
for r in range(13, 55):
    rows_by_name.setdefault(nm(str(ws[f'A{r}'].value or '')), []).append(r)

BODY_FONT = Font(name='Aptos Narrow', size=16)
AL = Alignment(horizontal='center', vertical='center', wrap_text=True)
# largura útil aproximada (caracteres por linha em fonte 16) por coluna
def chars_per_line(col):
    w = ws.column_dimensions[col].width if ws.column_dimensions[col].width else None
    # colunas agrupadas herdam a largura da primeira do grupo
    if not w or col not in ws.column_dimensions:
        for k, dim in ws.column_dimensions.items():
            if dim.min and dim.max and dim.min <= openpyxl.utils.column_index_from_string(col) <= dim.max:
                w = dim.width
    w = w or 35.7
    return max(10, int(w * 11 / 16 * 1.25))

def write_row(r, spec):
    lines_max = 0
    for key, items in spec.items():
        col = COL[key]
        txt = '\n\n'.join(items)
        c = ws[f'{col}{r}']
        c.value = txt
        c.font = BODY_FONT
        c.alignment = AL
        cpl = chars_per_line(col)
        lines = sum(max(1, math.ceil(len(p) / cpl)) for p in txt.split('\n'))
        lines_max = max(lines_max, lines)
    ws.row_dimensions[r].height = max(125.25, min(409, lines_max * 21 + 20))

usage = []  # (ambiente na tabela, linha(s) na grade)
dml_rows = rows_by_name['DML']
for amb, targets in MAP:
    spec = classify(R[amb])
    for t in targets:
        rr = rows_by_name[t]
        if t == 'DML':
            rr = dml_rows  # a tabela traz um único DML: aplicado aos dois DMLs do pilotis
        for r in rr:
            write_row(r, spec)
    usage.append((amb, ', '.join(targets)))

# novas linhas após a última do modelo (linha 54), com o estilo da linha 54
last = 54
ws.row_dimensions[56].height = None
ws['S61'].value = None
for i, (amb, name) in enumerate(NEW):
    r = last + 1 + i
    for c in range(1, 27):
        src = ws.cell(last, c); dst = ws.cell(r, c)
        dst.font = copy.copy(src.font); dst.fill = copy.copy(src.fill)
        dst.border = copy.copy(src.border); dst.alignment = copy.copy(src.alignment)
        dst.number_format = src.number_format
    ws.cell(r, 1).value = name
    ws.row_dimensions[r].height = 125.25
    write_row(r, classify(R[amb]))
    usage.append((amb, name + ' (linha nova)'))
end = last + len(NEW)
ws.row_dimensions[end + 1].height = 2.1
ws.print_area = f'A1:AF{end}'

# ambientes do modelo sem informação na tabela
filled = set()
for r in range(13, end + 1):
    if any(ws.cell(r, c).value for c in range(2, 27)): filled.add(r)
sem_info = [nm(str(ws[f'A{r}'].value)) for r in range(13, 55) if r not in filled]

# nomes na tabela para exibição
def disp(a):
    return {'VEST_TERMAS': 'Vestiário Fem| masc. - Thermas (piso P10)',
            'VEST_PISCINA': 'Vestiário Fem| masc. - Thermas (piso P7)'}.get(a, a)

# ---------- cabeçalhos (data / revisão) ----------
for w in wb.worksheets:
    for addr in ('A5',):
        if w[addr].value == 'XX/XX/XXXX': w[addr].value = DATA
    if w['A7'].value == 'REV. R0X': w['A7'].value = REV

# TÉRREO e 2º PAVTO: sem informação na tabela v00
for sn in ('TÉRREO ', '2° PAVTO'):
    w = wb[sn]
    w['B12'].value = 'GERAL — SEM ESPECIFICAÇÕES NA TABELA DE ACABAMENTOS V00 (A TABELA ABRANGE SOMENTE O 3º PAVIMENTO - PILOTIS)'

# ---------- controle de revisões ----------
cr = wb['CONTROLE DE REVISÕES ']
linhas = [
    (1, DATA, 'R00', 'PILOTIS (3º PAVIMENTO)', f'Emissão inicial. Especificações levantadas da {FONTE}.'),
    (2, DATA, 'R00', 'TÉRREO E 2º PAVIMENTO', 'Sem especificações na tabela v00 (abrange somente o 3º pavimento - Pilotis). Aguardando informação.'),
    (3, DATA, 'R00', 'IMAGENS ESPECIFICAÇÕES', 'Imagens das fichas TERMAS e VESTIÁRIOS - TERMAS da tabela v00 (págs. 7 e 8).'),
    (4, DATA, 'R00', 'PENDÊNCIAS', 'Nova aba com as pendências de especificação e a correspondência ambiente da tabela × linha da grade.'),
]
for i, row in enumerate(linhas):
    for j, v in enumerate(row):
        c = cr.cell(26 + i, 1 + j); c.value = v
        c.alignment = Alignment(horizontal='center' if j != 4 else 'left', vertical='center', wrap_text=True)

# ---------- imagens ----------
wi = wb['IMAGENS ESPECIFICAÇÕES ']
ICOL = dict(PISO='B', FILETE='C', PEITORIL='D', RODAPE='E', PAREDE='F', BANCADA='G', LOUCAS='I', VALVULAS='J',
            SIFOES='K', TORNEIRAS='L', BACIAS='M', DUCHAS='N', REGISTROS='O', ACION_CHUV='P', RALOS='Q',
            ELETRICOS='R', MACANETAS='S', PORTAS='T')
IMGS = [  # (linha, coluna, [xrefs], legenda)
    (13, 'PISO', [26], 'ESPAÇO TERMAS\nP10 - Piso elevado em mármore branco selecionado'),
    (13, 'PAREDE', [28, 30, 31], 'ESPAÇO TERMAS\nPA49 - Mosaico Bisazza (painel): New Sterlizia / Narciso'),
    (13, 'ACION_CHUV', [35, 37], 'ESPAÇO TERMAS\nD54 - Chuveiro de embutir no teto Raindream, Roca'),
    (14, 'PAREDE', [33, 30], 'ESPAÇO TERMAS\nPA49 - Mosaico Bisazza (pilar): Narciso'),
    (15, 'PISO', [48], 'VESTIÁRIO FEM./MASC. - TERMAS\nP10 - Piso elevado em mármore branco selecionado'),
    (15, 'PAREDE', [49], 'VESTIÁRIO FEM./MASC. - TERMAS\nD38 - Espelho liso incolor 4mm'),
    (15, 'BANCADA', [48], 'VESTIÁRIO FEM./MASC. - TERMAS\nD73 - Bancada em mármore branco selecionado'),
    (15, 'LOUCAS', [61], 'VESTIÁRIO FEM./MASC. - TERMAS\nD72 - Cuba de semi-encaixe slim, Deca'),
    (15, 'TORNEIRAS', [63], 'VESTIÁRIO FEM./MASC. - TERMAS\nD69 - Torneira com sensor Decalux, Deca'),
    (15, 'BACIAS', [55, 57, 59], 'VESTIÁRIO FEM./MASC. - TERMAS\nD22 - Bacia Carrara (Deca) / Panache (Kohler) e assento'),
    (15, 'ACION_CHUV', [69, 71], 'VESTIÁRIO FEM./MASC. - TERMAS\nD33 - Monocomando para chuveiro Mix&Match, Docol'),
    (16, 'PAREDE', [51], 'VESTIÁRIO FEM./MASC. - TERMAS\nPA47 - Pintura com tinta mineral, cor 930, Terracor'),
    (16, 'ACION_CHUV', [65, 67], 'VESTIÁRIO FEM./MASC. - TERMAS\nD70 - Chuveiro 200 de parede Docol Eden'),
]
UP = os.path.join(S, 'imagens_usuario')
HIRES = {28: f'{UP}/1.webp', 30: f'{UP}/2.webp', 31: f'{UP}/3.webp', 33: f'{UP}/4.webp', 65: f'{UP}/6.png', 67: f'{UP}/7.png', 69: f'{UP}/8.png', 71: f'{UP}/9.png'}
CAP_FONT = Font(name='Aptos Narrow', size=14, bold=False)
for r, key, xs, cap in IMGS:
    col = ICOL[key]
    c = wi[f'{col}{r}']
    c.value = cap
    c.font = CAP_FONT
    c.alignment = Alignment(horizontal='center', vertical='bottom', wrap_text=True)
    x_off = 8
    for x in xs:
        hi = HIRES.get(x)
        if hi:  # imagem em alta resolução enviada pelo usuário
            im = PILImage.open(hi).convert('RGB')
            if im.height > 900: im = im.resize((int(im.width * 900 / im.height), 900), PILImage.LANCZOS)
            p = f'{TMP}/hi_{r}_{col}_{x}.jpg'; im.save(p, quality=85)
            big = im.resize((int(im.width * 150 / im.height), 150))  # só para o tamanho de exibição
        else:
            im = PILImage.open(f'{S}/imagens_tabela/x{x}.png').convert('RGB')
            scale = 150 / im.height
            big = im.resize((int(im.width * scale), 150), PILImage.LANCZOS)
            p = f'{TMP}/big_{r}_{col}_{x}.png'; big.save(p)
        xi = XLImage(p)
        ci = openpyxl.utils.column_index_from_string(col) - 1
        marker = AnchorMarker(col=ci, colOff=pixels_to_EMU(x_off), row=r - 1, rowOff=pixels_to_EMU(10))
        xi.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(big.width), pixels_to_EMU(big.height)))
        wi.add_image(xi)
        x_off += big.width + 6

# ---------- aba de pendências ----------
wp = wb.create_sheet('PENDÊNCIAS')
hdr_font = Font(name='Aptos Narrow', size=12, bold=True, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='3A7D22')
title_font = Font(name='Aptos Narrow', size=16, bold=True)
body = Font(name='Aptos Narrow', size=11)
thin = Side(style='thin', color='808080'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
wp['A1'] = 'GRADE DE ESPECIFICAÇÕES PARA ÁREA COMUM - EDIFÍCIO JARDIM · PENDÊNCIAS E RASTREABILIDADE'
wp['A1'].font = title_font
wp['A2'] = f'Fonte: {FONTE}. Emissão {DATA}, {REV}.'
wp['A2'].font = body
for col, w in zip('ABCDE', (6, 34, 30, 70, 28)):
    wp.column_dimensions[col].width = w

def table(r0, title, header, rows):
    wp.cell(r0, 1, title).font = Font(name='Aptos Narrow', size=13, bold=True)
    for j, h in enumerate(header):
        c = wp.cell(r0 + 1, 1 + j, h); c.font = hdr_font; c.fill = hdr_fill; c.border = B
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = wp.cell(r0 + 2 + i, 1 + j, v); c.font = body; c.border = B
            c.alignment = Alignment(vertical='top', wrap_text=True, horizontal='center' if j == 0 else 'left')
    return r0 + 2 + len(rows) + 2

pend = [
    ('ESPAÇO TERMAS / VESTIÁRIO TERMAS / SAUNA A VAPOR / MASSAGEM / BELEZA', 'P10, PA46, PA48, PA23, D73, D59-D61, T10', 'Especificar qual o mármore branco ("mármore branco selecionado").', 'Tabela v00, págs. 7 e 8'),
    ('ESPAÇO TERMAS', 'PA49 - Mosaico New Sterlizia', 'Especificar o rejunte do mosaico New Sterlizia (a Bisazza recomenda Fillgel Plus 6601 rosa perla; a tabela indica 1101 Bianco Alabastro).', 'Tabela v00, pág. 7'),
    ('ESPAÇO TERMAS', 'PA48 - Mármore frisado', 'Especificar a espessura da pedra ("Espessura XX cm").', 'Tabela v00, pág. 7'),
    ('ESPAÇO TERMAS', 'T7 - Forro', 'Definir a especificação do forro resistente à umidade ("a ser definida").', 'Tabela v00, págs. 2 e 7'),
    ('VESTIÁRIO FEM./MASC. - TERMAS', 'D38 - Espelho', 'Confirmar se o espelho é colado sobre "armário em MDF" ou somente sobre chapa de MDF.', 'Tabela v00, pág. 8'),
    ('I.S / I.S P.C.D / I.S FEM. E MASC. / VESTIÁRIOS', 'D22 - Bacia sanitária', 'Definir entre Deca Carrara P606.17 (400x375x665mm) e Kohler Panache K-18395BR-0 (410x370x635mm). Informar se será entregue com assento (compatíveis: Deca AP.236.17/AP.237.17/AP.65.17; Kohler Puretide K-76923BR-0).', 'Tabela v00, pág. 8'),
    ('VESTIÁRIO FEM./MASC. (piso P7)', 'D31 - Chuveiro 200 Docol Eden', 'O código informado (90009312006) é o mesmo do monocomando D33; no outro vestiário o chuveiro D70 tem código 90008879006. Confirmar o código.', 'Tabela v00, pág. 2'),
    ('VESTIÁRIOS (P7 e P10)', 'Nome do ambiente', 'Na tabela os dois vestiários aparecem como "Vestiário Fem| masc. - Thermas" (págs. 1-3) e a repetição (págs. 4-6) chama o de piso P10 de "Vestiário Fem| masc." e o de piso P7 de "- Thermas". A grade segue a ficha da pág. 8: piso P10 = VESTIÁRIO FEM./MASC. - TERMAS (linha nova); piso P7 = VESTIÁRIO FEMININO / MASCULINO do modelo. Confirmar.', 'Tabela v00, págs. 2, 5 e 8'),
    ('VESTIÁRIO FEM./MASC. (piso P7)', 'Rodapé', 'Rodapé não informado na tabela (célula vazia).', 'Tabela v00, pág. 2'),
    ('MASSAGEM / BELEZA', 'D53 / D51 - Torneira Mix&Match Presence', 'Fabricante grafado "Socol" (provável Docol) e, no D51, "cor sensor" e "Mix@Match". Confirmar.', 'Tabela v00, pág. 2'),
    ('ESPAÇO DA BELEZA', 'Louças', 'Há torneira (D51) mas nenhuma cuba especificada. Confirmar.', 'Tabela v00, pág. 2'),
    ('QUADRA INFANTIL / QUADRA DE TÊNIS', 'Pinturas', 'Tons "a definir" (piso epóxi verde e cinza, paredes cinza e verde).', 'Tabela v00, págs. 2 e 3'),
    ('APOIO QUADRAS', 'P27 - Deck Tauari Teak', 'Dimensão das réguas grafada "200mmx100mmx2000m-acima". Confirmar.', 'Tabela v00, pág. 3'),
    ('DEPÓSITO / DEPÓSITO JARDIM', '12 - Porcelanato', 'Referência grafada "Khali pff white". Considerado Khali off white AC (Eliane).', 'Tabela v00, pág. 2'),
    ('DML', 'Todos os itens', 'A tabela traz um único DML; a especificação foi aplicada aos dois DMLs do pilotis do modelo. Confirmar.', 'Tabela v00, pág. 1'),
    ('LOBBY / I.S / ESPAÇO KIDS / SALÃO DE FESTAS / DEMAIS', 'Soleira, filete, peitoril, sóculo, válvulas, sifões, registros, elétricos, maçanetas, portas, janelas', 'Colunas do modelo sem correspondência na tabela de acabamentos (não informadas).', 'Modelo CDJ-GRE-ARQ-AC'),
    ('TÉRREO e 2º PAVIMENTO', 'Todos', 'A tabela v00 só traz o 3º pavimento - Pilotis. Abas mantidas sem preenchimento.', 'Tabela v00'),
]
r = table(4, '1. PENDÊNCIAS DE ESPECIFICAÇÃO', ['Nº', 'AMBIENTE', 'ITEM', 'PENDÊNCIA', 'ORIGEM'],
          [(i + 1,) + p for i, p in enumerate(pend)])
r = table(r, '2. AMBIENTES DO MODELO (PILOTIS) SEM INFORMAÇÃO NA TABELA', ['Nº', 'AMBIENTE NO MODELO', '', '', ''],
          [(i + 1, a, '', '', '') for i, a in enumerate(sem_info)])
r = table(r, '3. CORRESPONDÊNCIA: AMBIENTE NA TABELA × LINHA DA GRADE (PILOTIS)', ['Nº', 'AMBIENTE NA TABELA V00', 'LINHA(S) NA GRADE', '', ''],
          [(i + 1, disp(a), t, '', '') for i, (a, t) in enumerate(usage)])
r = table(r, '4. CRITÉRIO DE DISTRIBUIÇÃO NAS COLUNAS DA GRADE', ['Nº', 'COLUNA DA TABELA', 'COLUNA DA GRADE', 'CRITÉRIO', ''], [
    (1, 'PISO / RODAPÉ / TETO', 'PISO/REJUNTE / RODAPÉ / TETO', 'Transcrição direta.', ''),
    (2, 'PAREDE', 'PAREDE/REJUNTE', 'Revestimentos e pinturas (códigos PA, 01, 06, 09, 12...).', ''),
    (3, 'PAREDE', 'OUTROS REVESTIMENTOS', 'Espelhos sobre armário, fechamentos e painéis decorativos (códigos D), perfis de transição.', ''),
    (4, 'PAREDE', 'PORTAS', 'P28 (D43) - painel e porta em alumínio e vidro (Massagem).', ''),
    (5, 'PEDRAS', 'BANCADA / OUTROS REVESTIMENTOS', 'Bancadas em BANCADA; prateleira e testeira em OUTROS REVESTIMENTOS.', ''),
    (6, 'LOUÇAS E METAIS', 'LOUÇAS, CUBAS E TANQUES / TORNEIRAS / BACIAS / ACIONAMENTO DE CHUVEIRO / OUTROS', 'Separado por item; chuveiros em OUTROS LOUÇAS E METAIS.', ''),
    (7, 'DIVERSOS', 'RALOS / OUTROS COMPLEMENTARES', 'Ralos em RALOS; saboneteira, porta-papel e papeleira em OUTROS COMPLEMENTARES.', ''),
])
wp.sheet_view.showGridLines = False
wp.page_setup.orientation = 'landscape'
wp.page_setup.fitToWidth = 1; wp.page_setup.fitToHeight = 0
wp.sheet_properties.pageSetUpPr.fitToPage = True

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('ok', OUT, 'novas linhas até', end, 'sem info:', sem_info)
