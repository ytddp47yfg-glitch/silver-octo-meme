"""Conferência visual (2026-10-02) de todas as etiquetas das plantas de arquitetura PR003-PR013.

A leitura automática (arq_ambientes.py) só pegava grupos com 2+ etiquetas e ignorava as pranchas
PR008 e PR012 (trecho 2 do 2º pav. e do Pilotis). Cada ambiente foi conferido no desenho; o que
faltava na grade entra aqui. Quadro de especificações igual em todas as pranchas.

Uso: python3 arq_conferencia.py <grade.xlsx>
"""
import sys, copy, openpyxl
from openpyxl.styles import Font

XLSX = sys.argv[1]
Q07 = ('ARQ-P07 - Piso assentado em granito cinza santa rosa, acabamento serrado, Espessura: 2cm. Dimensões e '
       'paginação conforme detalhamento. Amostra a ser aprovada pela GPA&A. Obs.: Resina, hidrofugante ou outros '
       'acabamentos aplicados no granito devem ser foscos e não podem alterar ou escurecer a cor original do material.')
Q08 = ('ARQ-P08 - Piso elevado em granito cinza santa rosa, acabamento serrado, Espessura: 2cm. Dimensões e '
       'paginação conforme detalhamento. Amostra a ser aprovada pela GPA&A. Obs.: Resina, hidrofugante ou outros '
       'acabamentos aplicados no granito devem ser foscos e não podem alterar ou escurecer a cor original do material.')
Q21 = ('ARQ-P21 - Piso em granito marrom absoluto serrado e envernizado, ou pedra sintética a definir. '
       'Paginação conforme detalhamento. Amostra a ser aprovada pela GPA&A.')

PREENCHER = [  # (aba, ambiente, coluna, texto)
 ('TÉRREO ', 'RAMPA ACESSO PEDESTRES', 'B', Q21 + ' (Planta PR003: plano inclinado - pedestres (social), i=4,99%)'),
 ('PILOTIS', 'VARANDA', 'B', Q08 + ' (Plantas PR011 e PR012)'),
 # usuária (2026-10-02): lançar as circulações externas como dois ambientes
 ('PILOTIS', 'CIRCULAÇÃO JARDINS', 'B', Q08 + ' (Plantas PR011 e PR013: circulações externas junto aos jardins)'),
 ('PILOTIS', 'CIRCULAÇÃO PLAYGROUND', 'B', Q08 + ' (Planta PR013: circulação junto ao playground)'),
]
NOVAS = [  # linhas novas no fim da aba
 ('TÉRREO ', 'PLANO INCLINADO - PEDESTRES (SERVIÇO)', 'B', Q07 + ' (Planta PR003, i=3,83%)'),
 ('TÉRREO ', 'PLANO INCLINADO - VEÍCULOS', 'B', Q07 + ' (Planta PR003, dois trechos, i=3,50% e 3,28%)'),
]
NOTAS = [
 ('Conferência visual', 'Todos os ambientes das plantas PR003, PR004, PR005, PR007, PR008, PR009, PR011, PR012 e PR013 foram conferidos no desenho, etiqueta por etiqueta. Os ambientes já lançados estão de acordo com as plantas. Lista completa: levantamento/arquitetura/conferencia_etiquetas_arquitetura.csv.'),
 ('PR008 e PR012', 'São o trecho 2 do 2º pavimento e do Pilotis (não revisões antigas, como estava anotado). PR008 só traz a rampa (○04 △07 □03, igual à PR009). PR012 traz a varanda (○08); piscinas, prainha e deck sem etiqueta.'),
 ('PILOTIS – VARANDA', 'Planta indica só piso ○08 (ARQ-P08), lançado na grade. Detalhamento R04 sem especificação.'),
 ('PILOTIS – CIRCULAÇÕES EXTERNAS', 'As circulações de 40,36 m² (PR011), 59,22 m² e 83,52 m² (PR013) têm só piso ○08 (ARQ-P08), a mesma especificação já lançada em "CIRCULAÇÃO EXTERNA COBERTA/DESCOBERTA E VARANDA ESPAÇO KIDS". Lançadas como dois ambientes, "CIRCULAÇÃO JARDINS" e "CIRCULAÇÃO PLAYGROUND", conforme orientação de 02/10/2026.'),
 ('TÉRREO – PLANOS INCLINADOS', 'Plano inclinado de pedestres (social) com ○21, lançado em RAMPA ACESSO PEDESTRES. Plano inclinado de pedestres (serviço) e de veículos com ○07, em linhas novas.'),
 ('TÉRREO – MUROS', 'Triângulo △04 isolado em três pontos do muro junto ao jardim (PR003): parede em textura acrílica fosca Terra Fértil, Suvinil. Não é um ambiente da grade.'),
 ('Sem etiqueta na planta', 'Elevadores, escadas, I.S. P.C.D. e vestíbulo (Térreo), DML e jardinagem (Térreo), jardins sobre laje, piscinas, prainha, deck, playground, lounge quadras, saunas, apoio e DML do salão, lobby pilotis e circulações internas do Pilotis. Hall, lobby, antecâmara (eclusa), espaço kids, gourmet, salão, cozinha, I.S., termas, academia, bar e vestiários do Pilotis têm "ACABAMENTOS VER INTERIORES".'),
]

wb = openpyxl.load_workbook(XLSX)
def linha(ws, nome):
    return next((r for r in range(13, ws.max_row + 1) if (ws.cell(r, 1).value or '').strip() == nome), None)
for aba, amb, col, txt in PREENCHER:
    ws = wb[aba]; r = linha(ws, amb)
    ws[f'{col}{r}'].value = txt
for aba, amb, col, txt in NOVAS:
    ws = wb[aba]; r = linha(ws, amb)
    if r is None:
        last = max(x for x in range(13, ws.max_row + 1) if ws.cell(x, 1).value)
        r = last + 1
        for c in range(1, 27):
            ws.cell(r, c)._style = copy.copy(ws.cell(last, c)._style); ws.cell(r, c).value = None
        ws.cell(r, 1).value = amb
    ws[f'{col}{r}'].value = txt
p = wb['PENDÊNCIAS']
if not any(p.cell(r, 1).value and str(p.cell(r, 1).value).startswith('7. CONFERÊNCIA') for r in range(1, p.max_row + 1)):
    r = p.max_row + 2
    p.cell(r, 1, '7. CONFERÊNCIA VISUAL DAS PLANTAS DE ARQUITETURA (02/10/2026)').font = Font(bold=True, size=12)
    for k, v in NOTAS:
        r += 1; p.cell(r, 2, k); p.cell(r, 4, v)
        p.cell(r, 4).alignment = copy.copy(p.cell(r - 1, 4).alignment)
for r in range(1, p.max_row + 1):
    v = p.cell(r, 4).value
    if v == 'Pranchas de revisão anterior (R11 e R13 de 2025) às demais; não foram usadas.':
        p.cell(r, 4).value = 'Trecho 2 do 2º pavimento e do Pilotis; conferidas em 02/10/2026 (ver item 7).'
wb.save(XLSX)
print('ok')
