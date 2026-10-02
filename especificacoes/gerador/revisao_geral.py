"""Revisão geral da grade (2026-10-02): textos, notas e sumário desatualizados.

Uso: python3 revisao_geral.py <grade.xlsx>
"""
import sys, re, copy, openpyxl
from openpyxl.styles import Font

XLSX = sys.argv[1]
ABAS = ['TÉRREO ', '2° PAVTO', 'PILOTIS', 'APTO TIPO']
NOMES = {'HALL GARGEM': 'HALL GARAGEM', 'CIRCULAÇÃO ELEVADOR DE SERVIÇO/ANTICÂRA': 'CIRCULAÇÃO ELEVADOR DE SERVIÇO/ANTECÂMARA',
         'PISCINA DECOBERTA': 'PISCINA DESCOBERTA'}
TEXTO = [(r'caravalho', 'carvalho'), (r'Khali pff white', 'Khali off white'), (r'Scoth Brite', 'Scotch Brite'),
         (r'hidrofulgante', 'hidrofugante'), (r'(\d)\.\. Ref\.', r'\1. Ref.'), (r'Ref\.:(?=[A-Za-z])', 'Ref.: ')]

def limpa(v):
    if not isinstance(v, str): return v
    for a, b in TEXTO: v = re.sub(a, b, v)
    return '\n'.join(re.sub(r'(?<=\S) {2,}', ' ', l).rstrip() for l in v.split('\n'))

GERAL = {'TÉRREO ': 'GERAL — TÉRREO PELO DETALHAMENTO DE INTERIORES R04 E PELAS PLANTAS DE ARQUITETURA PR003 A PR005',
         '2° PAVTO': 'GERAL — 2º PAVIMENTO PELAS PLANTAS DE ARQUITETURA PR007 A PR009 (O DETALHAMENTO R04 SÓ TRAZ MOBILIÁRIO)'}

wb = openpyxl.load_workbook(XLSX)
for n in ABAS:
    ws = wb[n]
    if n in GERAL: ws['B12'].value = GERAL[n]
    for r in range(13, ws.max_row + 1):
        a = ws.cell(r, 1)
        if isinstance(a.value, str):
            a.value = NOMES.get(a.value.strip(), a.value.strip())
        for c in range(2, 27):
            cel = ws.cell(r, c)
            cel.value = limpa(cel.value)
            cor = cel.font.color
            if cor is not None and cor.type == 'theme' and cor.theme == 5:  # laranja de INFO. REVISADA que sobrou
                f = cel.font; cel.font = Font(name=f.name, size=f.sz, bold=f.b, italic=f.i)

# CONTROLE DE REVISÕES: sumário por aba e registro da R00
ct = wb['CONTROLE DE REVISÕES ']
ct['A11'].value = 'ABA'
SUM = [('CONTROLE DE REVISÕES', 'SUMÁRIO E CONTROLE DE REVISÕES'),
       ('TÉRREO', 'ESPECIFICAÇÃO DE REVESTIMENTOS, LOUÇAS E METAIS E COMPLEMENTARES'),
       ('2° PAVTO', 'ESPECIFICAÇÃO DE REVESTIMENTOS, LOUÇAS E METAIS E COMPLEMENTARES'),
       ('PILOTIS', 'ESPECIFICAÇÃO DE REVESTIMENTOS, LOUÇAS E METAIS E COMPLEMENTARES; IMAGENS AO LADO DA GRADE (LINK "VER IMAGEM")'),
       ('APTO TIPO', 'ESPECIFICAÇÃO DO APARTAMENTO TIPO'),
       ('CONFLITOS', 'ESPECIFICAÇÕES CONFLITANTES (ARQUITETURA × DETALHAMENTO; PRANCHA × TABELA DA PRÓPRIA PRANCHA)'),
       ('PENDÊNCIAS', 'PENDÊNCIAS E RASTREABILIDADE DAS FONTES')]
for i in range(11):
    r = 12 + i
    a, c = SUM[i] if i < len(SUM) else (None, None)
    ct.cell(r, 1).value = a; ct.cell(r, 3).value = c
LOG = [('TODAS', 'Emissão inicial (primeira versão). Prevalece o detalhamento sobre a Tabela 01; divergências entre arquitetura e detalhamento, e entre prancha e tabela da própria prancha, ficam em vermelho e na aba CONFLITOS.'),
       ('PILOTIS (3º PAVIMENTO)', 'Tabela de acabamentos Ed. Cidade Jardim v00 (19/06/2026), detalhamento de interiores R04 (23/04/2026) e plantas de arquitetura PR011 a PR013.'),
       ('TÉRREO E 2º PAVIMENTO', 'Detalhamento de interiores R04 (23/04/2026) e plantas de arquitetura CDJ-ARQ-EX PR003 a PR009, conferidas no desenho etiqueta por etiqueta.'),
       ('APTO TIPO', 'Tabela JARDIM-ACABAMENTOS R01 e detalhamento do apartamento (CDJ_DET_EX PR001-PR008, CDJ-ARQ-DT PR406/407/408).'),
       ('IMAGENS', 'Imagens das fichas TERMAS e VESTIÁRIOS - TERMAS na aba PILOTIS, à direita da grade (coluna AH em diante): "► VER IMAGEM (clique)" leva à foto e "◄ VOLTAR" retorna.')]
MODELO = [copy.copy(ct.cell(29, c)._style) for c in range(1, 10)]  # linha APTO TIPO: fonte preta, tamanho padrão
for i, (amb, com) in enumerate(LOG):
    r = 26 + i
    for c in range(1, 10): ct.cell(r, c)._style = copy.copy(MODELO[c - 1])
    ct.row_dimensions[r].height = ct.row_dimensions[29].height
    ct.cell(r, 1).value = str(i + 1); ct.cell(r, 2).value = '30/09/2026'; ct.cell(r, 3).value = 'R00'
    ct.cell(r, 4).value = amb; ct.cell(r, 5).value = com

# PENDÊNCIAS
p = wb['PENDÊNCIAS']
p['A2'].value = ('Fontes: Tabela de acabamentos Ed. Cidade Jardim v00 (19/06/2026); detalhamento de interiores R04 (23/04/2026); '
                 'plantas de arquitetura CDJ-ARQ-EX PR003–PR013; apartamento tipo: tabela JARDIM-ACABAMENTOS R01 e detalhamento. '
                 'Emissão 30/09/2026, REV. R00.')
heads = [r for r in range(3, p.max_row + 1) if isinstance(p.cell(r, 1).value, str) and re.match(r'^\d+\. [A-ZÁÉÍÓÚÂÊÔÃÕÇ]', p.cell(r, 1).value)]
for k, r in enumerate(heads, 1):
    p.cell(r, 1).value = re.sub(r'^\d+\.', f'{k}.', p.cell(r, 1).value)
num = {re.sub(r'^\d+\. ', '', p.cell(r, 1).value)[:20]: i for i, r in enumerate(heads, 1)}
for r in range(1, p.max_row + 1):
    if p.cell(r, 4).value == 'A tabela v00 só traz o 3º pavimento - Pilotis. Abas mantidas sem preenchimento.':
        p.cell(r, 4).value = 'A tabela v00 só traz o 3º pavimento - Pilotis. As abas foram preenchidas pelo detalhamento R04 e pelas plantas de arquitetura (ver itens 5 e 7).'
        p.cell(r, 5).value = 'Tabela v00 → R04 e arquitetura'
# item 2: situação atual de cada ambiente
pil = wb['PILOTIS']
cheio = {}
for r in range(13, pil.max_row + 1):
    a = (pil.cell(r, 1).value or '').strip()
    cheio[a] = cheio.get(a, False) or any(pil.cell(r, c).value for c in range(2, 27))
r2 = heads[1]
p.cell(r2, 1).value = '2. AMBIENTES DO MODELO (PILOTIS) SEM INFORMAÇÃO NA TABELA V00 — SITUAÇÃO ATUAL'
p.cell(r2 + 1, 4).value = 'SITUAÇÃO'
p.cell(r2 + 1, 4)._style = copy.copy(p.cell(r2 + 1, 2)._style)
for r in range(r2 + 2, heads[2]):
    a = (p.cell(r, 2).value or '').strip()
    if not a or (p.cell(r, 4).value or '').startswith('Tabela v00 e detalhamento'): continue
    p.cell(r, 4).value = ('Preenchido pelo detalhamento R04 e/ou pelas plantas de arquitetura.' if cheio.get(a)
                          else 'Sem especificação em nenhuma fonte recebida (tabela v00, detalhamento R04 e plantas de arquitetura sem etiqueta). Definir.')
def nomes(v):
    for a, b in NOMES.items(): v = v.replace(a, b)
    return v
for aba in ('PENDÊNCIAS', 'CONFLITOS'):
    for row in wb[aba].iter_rows():
        for c in row:
            if isinstance(c.value, str): c.value = nomes(limpa(c.value))
wb.save(XLSX)
print('ok', len(heads), 'seções')
