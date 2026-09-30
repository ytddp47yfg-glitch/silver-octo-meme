"""Atualiza a grade CDJ-GRE-ARQ-AC-R00 com a lista de especificações de interiores R04
(CDJ-AIN-EX-DC001-ESPECIFIC-R04) e as pranchas de detalhamento CDJ-AIN-EX-PR001..PR042 R04.

Uso: python3 levantamento_r04.py <interiores_R04.json> <grade.xlsx>
"""
import json, re, sys, copy
import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.styles.colors import Color

SRC, XLSX = sys.argv[1], sys.argv[2]
groups = json.load(open(SRC))
DESC = {}
for g in groups:
    for it in g['itens']:
        if re.fullmatch(r'(PI|P|PA|T|FO|D)\d+', it['cod']) and it['desc'] != '#N/D':
            DESC.setdefault(it['cod'], (it['desc'], it['sim']))

# colunas da grade (linha 11 das abas de pavimento)
COL = dict(piso='B', rodape='F', parede='H', teto='I', bancada='J', outrosrev='K', loucas='L',
           torneiras='O', bacias='P', duchas='Q', registros='R', chuveiro_ac='S', ralos='T',
           outros_lm='U', eletricos='V', portas='X', outros_c='Z')

def coluna(cod):
    d, s = DESC[cod]
    dl = d.lower()
    if s == '○': return 'piso'
    if s == '□': return 'teto'
    if s == '∆': return 'rodape' if cod in ('PA20', 'PA23') else 'parede'
    if cod == 'D12': return 'eletricos'
    if re.match(r'(bancada|mureta)', dl): return 'bancada'
    if re.match(r'(cuba|tanque|lavatório)', dl): return 'loucas'
    if re.match(r'(torneira|monocomando de mesa|misturador)', dl): return 'torneiras'
    if dl.startswith('bacia'): return 'bacias'
    if dl.startswith('ducha higiênica'): return 'duchas'
    if dl.startswith('acabamento para registro'): return 'registros'
    if dl.startswith('monocomando para chuveiro'): return 'chuveiro_ac'
    if dl.startswith('ralo'): return 'ralos'
    if dl.startswith('chuveiro'): return 'outros_lm'
    if re.match(r'(porta|divisórias e portas|painel e porta)', dl): return 'portas'
    if cod in ('D37', 'D38', 'D75'): return 'outrosrev'
    return 'outros_c'

# ambiente da grade -> códigos da lista R04 (divisão dos grupos conferida nas pranchas)
PILOTIS = {
 'LOBBY PILOTIS': 'PI02 P7 PA20 PA44 PA53 PA58 FO2 T1 D12 D19 D39 D74 D75 D85',
 'CIRCULAÇÃO': 'PI02 P7 PA20 PA44 PA53 PA58 FO2 T1 D12 D75',
 'SALÃO DE FESTAS': 'PI02 P7 PA20 PA40 PA41 PA31 FO1 FO2 T1 D12 D21 D24 D83',
 'I.S FEMININO': 'PI02 P7 PA20 PA39 T1 D12 D22 D23 D25 D26 D27 D28 D37 D83',
 'I.S MASCULINO': 'PI02 P7 PA20 PA39 T1 D12 D22 D23 D25 D26 D27 D28 D37 D83',
 'COZINHA SALÃO DE FESTAS': 'PI02 P7 PA43 PA52 T1 D12 D21 D36 D74 D80 D81 D82',
 'I.S P.C.D': 'P7 PA20 PA26 PA32 PA52 T1 D12 D19 D22 D23 D26 D27 D28 D30 D38',
 'I.S': 'P7 PA20 PA26 PA32 PA52 T1 D12 D19 D22 D23 D26 D27 D28 D30 D38',
 'ESPAÇO KIDS': 'PI02 P24 PA20 PA34 PA35 PA36 PA37 PA38 FO1 FO2 T1 D12 D18 D11 D39 D56',
 'ESPAÇO GOURMET': 'PI02 P7 PA20 PA44 PA32 PA33 FO1 FO2 D8 D9 D10 D11 D12 D13 D14 D15 D16 D17 D83',
 'PISCINA DECOBERTA': 'P18 P28 PA59 D76 D77 D78 D79',
 'DECK': 'P26',
 'ESTAR PISCINA / CIRCULAÇÃO': 'P28 P8 D76 D78',
 'BAR PISCINA': 'P9 PA20 PA42 PA51 T4 T1 D9 D10 D12 D34 D35 D46 D47 D48 D49 D50',
 'CIRCULAÇÃO VESTIÁRIOS': 'P9 PA52 T1 D12',
 'VESTIÁRIO FEMININO': 'P9 PA20 PA26 PA52 PA60 T1 D12 D22 D26 D27 D28 D29 D31 D32 D33 D34 D35 D36 D37',
 'VESTIÁRIO MASCULINO': 'P9 PA20 PA26 PA52 PA60 T1 D12 D22 D26 D27 D28 D29 D31 D32 D33 D34 D35 D36 D37',
 'ACADEMIA': 'PI02 P22 PA20 PA26 PA27 PA28 PA29 PA30 T1 FO2 D3 D4 D5 D6 D7 D12 D57',
 'PISCINA RAIA': 'P10 PA46',
 'SPA': 'P18 PA46',
 'ESPAÇO TERMAS': 'P10 PA23 PA46 PA48 PA49 T1 T7 D12 D35 D43 D54 D58',
 'SAUNA VAPOR': 'P10 PA46 T10 D35 D54 D58',
 'SAUNA SECA': 'P19 PA50 T11',
 'MASSAGEM': 'P10 PA23 PA26 PA46 PA47 T1 D12 D43 D51 D53 D55 D59 D60 D61 D62',
 'BELEZA': 'P10 PA23 PA26 PA46 PA47 T1 D12 D43 D51 D53 D55 D59 D60 D61 D62',
 'HALL VESTIÁRIOS': 'P10 PA23 PA46 PA48 T1 D39 D43',
 'VEST. P.C.D': 'P10 PA23 PA46 PA47 T8 D12 D22 D26 D27 D28 D33 D35 D38 D69 D70 D72 D73 D84',
 'DML#2': 'P10 PA23 PA47 T1 D67 D68',
 'VESTIÁRIO FEM./MASC. - TERMAS': 'P10 PA23 PA46 PA47 PA48 T8 D12 D22 D26 D27 D28 D33 D34 D35 D38 D39 D43 D69 D70 D72 D73',
 'APOIO QUADRAS': 'P7 P27 FO2 T1',
}
TERREO = {
 'ANTECÂMARA(ECLUSA PEDESTRE)': 'P13 PA20 PA56 T1 D12 D41 D42',
 'GUARITA': 'P5 PA1 PA20 T1 D12',
 'LOBBY ENTRADA': 'P13 PA1 PA20 PA53 PA54 PA56 T1 T12 D12 D41 D42 D63',
 'HALL ELEVADORES': 'P13 PA1 PA20 PA53 PA54 T1 D12 D64',
 'HALL GARGEM': 'P13 PA1 PA20 PA53 T1 D12 D44 D45 D64',
 'VESTÍBULO': 'P13 PA20 PA53 PA55 T1 D12 D39 D42',
 'I.S P.C.D': 'P13 PA20 PA26 PA53 T1 D12 D22 D26 D27 D28 D34 D37 D39 D65 D66',
}
INFERIDO = {  # divisões que as pranchas não deixam 100% claras
 'MASSAGEM': 'Massagem e Beleza estão juntas na lista R04 e nas pranchas 05 e 06; louças e bancadas listadas nos dois ambientes.',
 'BELEZA': 'Massagem e Beleza estão juntas na lista R04 e nas pranchas 05 e 06; louças e bancadas listadas nos dois ambientes.',
 'VEST. P.C.D': 'Vestiários P.C.D. das Termas não têm ampliação própria; itens seguem os vestiários e o DML das Termas (pranchas 10 e 11) e a ducha higiênica D84.',
 'PISCINA RAIA': 'Pranchas 01 e 04 indicam piscina revestida em pedra (mármore) com bordas abauladas.',
}

def texto(codes):
    por_col = {}
    for c in codes.split():
        por_col.setdefault(coluna(c), []).append(f'{c} - {DESC[c][0]}')
    return por_col

wb = openpyxl.load_workbook(XLSX)
ref = wb['PILOTIS']['K7']
REV_FILL = copy.copy(ref.fill); REV_FONT_COLOR = copy.copy(ref.font.color)
log = []

def aplica(ws, mapa, pav):
    vistos = {}
    for r in range(13, ws.max_row + 1):
        nome = (ws.cell(r, 1).value or '').strip()
        if not nome: continue
        vistos[nome] = vistos.get(nome, 0) + 1
        chave = nome if vistos[nome] == 1 or f'{nome}#{vistos[nome]}' not in mapa else f'{nome}#{vistos[nome]}'
        if nome == 'DML' and vistos[nome] == 1 and pav == 'PILOTIS':
            continue  # DML do salão: não consta na lista R04, fica com a v00
        if chave not in mapa: continue
        novo = texto(mapa[chave])
        for k, col in COL.items():
            cel = ws[f'{col}{r}']
            antigo = (cel.value or '').strip()
            val = '\n\n'.join(novo.get(k, []))
            ant_cod = set(re.findall(r'\b(?:PI|PA|FO|P|T|D)\d+\b', antigo))
            nov_cod = set(re.findall(r'^(?:PI|PA|FO|P|T|D)\d+', val, re.M))
            if not val and not antigo: continue
            if not val and antigo:
                continue  # informação da v00 que a R04 não cita: mantida
            if val != antigo:
                cel.value = val
                if ant_cod != nov_cod or not antigo:
                    cel.fill = copy.copy(REV_FILL)
                    f = copy.copy(cel.font); f.color = copy.copy(REV_FONT_COLOR); cel.font = f
                    log.append((pav, nome, ws[f'{col}11'].value.replace('\n', ' ').strip(),
                                ', '.join(sorted(ant_cod)) or '—', ', '.join(sorted(nov_cod))))
        ws.row_dimensions[r].height = None
    faltou = set(mapa) - {k if '#' not in k else k for k in mapa if k.split('#')[0] in vistos}
    return faltou

f1 = aplica(wb['PILOTIS'], PILOTIS, 'PILOTIS')
f2 = aplica(wb['TÉRREO '], TERREO, 'TÉRREO')
print('não encontrados:', f1, f2)

# cabeçalho: fonte
for ws in (wb['PILOTIS'], wb['TÉRREO ']):
    for r in range(1, 10):
        for c in ws[r]:
            if isinstance(c.value, str) and 'Tabela de acabamentos' in c.value:
                c.value = c.value + ' + Lista de especificações de interiores R04 (23/04/2026)'

# aba de pendências: novas linhas
p = wb['PENDÊNCIAS']
r = p.max_row + 2
p.cell(r, 1, '3. LEVANTAMENTO PELOS PROJETOS DE DETALHAMENTO R04 (23/04/2026)').font = Font(bold=True, size=12)
r += 1
for i, h in enumerate(['Nº', 'PAVIMENTO / AMBIENTE', 'COLUNA', 'ANTES (v00)', 'AGORA (R04)']):
    c = p.cell(r, 1 + i, h); c.font = Font(bold=True, color='FFFFFFFF'); c.fill = PatternFill('solid', fgColor='FF243F2E')
for n, (pav, nome, col, a, b) in enumerate(log, 1):
    r += 1
    for i, v in enumerate([n, f'{pav} – {nome}', col, a, b]): p.cell(r, 1 + i, v)
r += 2
p.cell(r, 1, '4. OBSERVAÇÕES DO LEVANTAMENTO R04').font = Font(bold=True, size=12)
obs = [(k, v) for k, v in INFERIDO.items()] + [
 ('VESTIÁRIO FEMININO / MASCULINO', 'Pranchas 37 e 38 citam torneira Argon (Docol) e bancada em granito Branco Siena; a lista R04 traz D29 (Decalux, Deca) e D36 (granito cinza Santa Rosa). Confirmar com a arquitetura.'),
 ('BAR PISCINA', 'Prancha 35 é imagem sem texto; não foi possível conferir.'),
 ('SALÃO DE FESTAS', 'Prancha 20 é imagem sem texto; divisão conferida nas pranchas 21 a 24.'),
 ('TÉRREO', 'A lista R04 agrupa antecâmara, guarita, lobby, vestíbulo, I.S., hall de elevadores e hall da garagem; a divisão por ambiente foi feita pelas pranchas 40 a 42.'),
 ('GOURMET EXTERNO / VARANDA / GUARITA PILOTIS', 'Grupos da lista R04 com códigos sem descrição (#N/D); não entraram na grade.'),
 ('2º PAVIMENTO', 'A lista R04 só traz mobiliário do 2º pavimento; aba continua sem especificação de acabamentos.'),
]
for k, v in obs:
    r += 1; p.cell(r, 2, k); p.cell(r, 4, v)
wb.save(XLSX)
json.dump({'pilotis': PILOTIS, 'terreo': TERREO, 'mudancas': log}, open(XLSX.replace('.xlsx', '_R04_log.json'), 'w'), ensure_ascii=False, indent=1)
print(len(log), 'células revisadas')
