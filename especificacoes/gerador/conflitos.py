"""Marca especificações conflitantes na grade: mantém as duas versões na célula e pinta de vermelho claro.

Uso: python3 conflitos.py <grade.xlsx>
"""
import sys, copy
import openpyxl
from openpyxl.styles import PatternFill, Font

XLSX = sys.argv[1]
FILL = PatternFill('solid', fgColor='FFFFC7CE')
COR = 'FF9C0006'

# (aba, ambiente, coluna, texto da segunda fonte)
CONFLITOS = [
 ('PILOTIS', 'VESTIÁRIO FEMININO', 'J', 'CONFLITO – Pranchas 37 e 38 do detalhamento R04: bancada em granito Branco Siena com duas cubas esculpidas com bandeja removível'),
 ('PILOTIS', 'VESTIÁRIO MASCULINO', 'J', 'CONFLITO – Pranchas 37 e 38 do detalhamento R04: bancada em granito Branco Siena com duas cubas esculpidas com bandeja removível'),
 ('PILOTIS', 'VESTIÁRIO FEMININO', 'O', 'CONFLITO – Pranchas 37 e 38 do detalhamento R04: torneira de mesa bica baixa Argon, Docol'),
 ('PILOTIS', 'VESTIÁRIO MASCULINO', 'O', 'CONFLITO – Pranchas 37 e 38 do detalhamento R04: torneira de mesa bica baixa Argon, Docol'),
]

wb = openpyxl.load_workbook(XLSX)

# legenda: INFO. CONFLITANTE na linha 9, abaixo de INFO. ANTERIOR / INFO. REVISADA
for nome in ('TÉRREO ', '2° PAVTO', 'PILOTIS'):
    ws = wb[nome]
    for a, b in (('J', 'K'), ('T', 'U'), ('AE', 'AF')):
        ref = ws[f'{b}7']
        if ref.value is None: continue
        rng = f'{a}9:{b}9'
        if rng not in [str(m) for m in ws.merged_cells.ranges]:
            ws.merge_cells(rng)
        c = ws[f'{a}9']
        c.value = 'INFO. CONFLITANTE (manter as duas)'
        c.fill = FILL
        f = copy.copy(ref.font); f.color = COR; c.font = f
        c.alignment = copy.copy(ref.alignment); c.border = copy.copy(ref.border)
        ws[f'{b}9'].border = copy.copy(ws[f'{b}7'].border)

for aba, amb, col, txt in CONFLITOS:
    ws = wb[aba]
    for r in range(13, ws.max_row + 1):
        if (ws.cell(r, 1).value or '').strip() != amb: continue
        cel = ws[f'{col}{r}']
        atual = cel.value or ''
        if txt not in atual:
            cel.value = (atual + '\n\n' + txt).strip()
        cel.fill = FILL
        f = copy.copy(cel.font); f.color = COR; cel.font = f
wb.save(XLSX)
print('ok')
