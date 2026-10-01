"""Área de impressão dinâmica em todas as abas da grade CDJ-GRE-ARQ-AC.

O Print_Area de cada aba vira uma fórmula OFFSET que vai até a última linha
com qualquer conteúdo nas colunas da aba, então linhas novas entram sozinhas
na impressão. Também grava uma área fixa equivalente antes, como reserva.
"""
import re, sys, zipfile, shutil, os
from html import escape
import openpyxl
from openpyxl.utils import get_column_letter

P = sys.argv[1] if len(sys.argv) > 1 else 'especificacoes/CDJ-GRE-ARQ-AC-R00.xlsx'
MAXR = 3000
# aba -> (última coluna, linha mínima, linhas de título)
CFG = {
    'CONTROLE DE REVISÕES ': ('I', 1, None),
    'TÉRREO ': ('AF', 11, '$1:$11'),
    '2° PAVTO': ('AF', 11, '$1:$11'),
    'PILOTIS': ('AF', 11, '$1:$11'),
    'APTO TIPO': ('AF', 11, '$1:$11'),
    'IMAGENS ESPECIFICAÇÕES ': ('U', 26, '$1:$11'),
    'PENDÊNCIAS': ('E', 1, None),
}

def ultima(ws, lc):
    n = openpyxl.utils.column_index_from_string(lc)
    return max((c.row for r in ws.iter_rows(max_col=n) for c in r
                if c.value is not None and str(c.value).strip() != ''), default=1)

wb = openpyxl.load_workbook(P)
dyn = {}
for i, ws in enumerate(wb.worksheets):
    lc, minr, tit = CFG[ws.title]
    ws.print_area = f'A1:{lc}{max(minr, ultima(ws, lc))}'
    if tit:
        ws.print_title_rows = tit
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    q = "'" + ws.title.replace("'", "''") + "'"
    rng = f"{q}!$A$1:${lc}${MAXR}"
    ncol = openpyxl.utils.column_index_from_string(lc)
    dyn[i] = (f"OFFSET({q}!$A$1,0,0,MAX({minr},MAX(INDEX(({rng}&\"\"<>\"\")*ROW({rng}),0))),{ncol})")
wb.save(P)

# troca a área fixa pela fórmula dinâmica no workbook.xml
tmp = P + '.tmp'
with zipfile.ZipFile(P) as zi, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        data = zi.read(it.filename)
        if it.filename == 'xl/workbook.xml':
            s = data.decode('utf8')
            def rep(m):
                return m.group(1) + escape(dyn[int(m.group(2))], quote=False) + m.group(3)
            s, k = re.subn(r'(<definedName name="_xlnm\.Print_Area" localSheetId="(\d+)"[^>]*>)[^<]*(</definedName>)', rep, s)
            assert k == len(dyn), k
            data = s.encode('utf8')
        zo.writestr(it, data)
shutil.move(tmp, P)
print('ok', {wb.worksheets[i].title: wb.worksheets[i].print_area for i in dyn})
