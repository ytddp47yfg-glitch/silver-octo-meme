"""Aba CONFLITOS: lista as células com especificações conflitantes (vermelho claro).

Lê as abas de pavimento, pega as células com preenchimento FFC7CE e monta uma
tabela com aba, ambiente, item, célula (com link) e o texto com as versões.
Recria a aba a cada execução. Rodar depois de imagens_links.py e antes de
area_impressao.py.
"""
import sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.hyperlink import Hyperlink

P = sys.argv[1] if len(sys.argv) > 1 else 'especificacoes/CDJ-GRE-ARQ-AC-R00.xlsx'
ABAS = ['TÉRREO ', '2° PAVTO', 'PILOTIS', 'APTO TIPO']
NOME = 'CONFLITOS'
VERDE, CINZA, VERM, VERM_TXT = '243F2E', 'E2E2E2', 'FFC7CE', '9C0006'
fino = Side(style='thin', color='8C8C8C'); BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)

wb = openpyxl.load_workbook(P)
for ws in wb.worksheets:   # preserva bytes das imagens existentes
    for im in ws._images:
        b = im._data(); im._data = (lambda b=b: b)

itens = []
for aba in ABAS:
    ws = wb[aba]
    for row in ws.iter_rows(min_row=13, max_col=26):
        for x in row:
            f = x.fill
            if x.value and f is not None and f.fill_type and f.fgColor.type == 'rgb' and (f.fgColor.rgb or '').endswith(VERM):
                item = str(ws.cell(11, x.column).value or '').strip()
                grupo = 'REVESTIMENTOS' if x.column <= 11 else ('LOUÇAS E METAIS' if x.column <= 21 else 'COMPLEMENTARES')
                texto = str(x.value).replace('\n\n► VER IMAGEM (clique)', '')
                itens.append((aba.strip(), str(ws.cell(x.row, 1).value or '').strip(), grupo, item, x.coordinate, aba, texto))

if NOME in wb.sheetnames:
    del wb[NOME]
pos = wb.sheetnames.index('PENDÊNCIAS') if 'PENDÊNCIAS' in wb.sheetnames else len(wb.sheetnames)
ws = wb.create_sheet(NOME, pos)
ws.sheet_view.showGridLines = False
larg = {'A': 6, 'B': 14, 'C': 34, 'D': 20, 'E': 22, 'F': 11, 'G': 90, 'H': 30}
for k, v in larg.items():
    ws.column_dimensions[k].width = v

ws.merge_cells('A1:H1')
t = ws['A1']; t.value = 'GRADE DE ESPECIFICAÇÕES · EDIFÍCIO JARDIM · ESPECIFICAÇÕES CONFLITANTES'
t.font = Font(name='Aptos Narrow', size=16, bold=True, color='FFFFFF'); t.fill = PatternFill('solid', fgColor=VERDE)
t.alignment = Alignment(horizontal='left', vertical='center', indent=1); ws.row_dimensions[1].height = 30
ws.merge_cells('A2:H2')
n = ws['A2']
n.value = ('Células em que projetos diferentes trazem especificações diferentes para o mesmo item. As duas versões '
           'foram mantidas na grade, em vermelho claro, até a definição da arquitetura. Clique na célula (coluna F) para ir até ela.')
n.font = Font(name='Aptos Narrow', size=11, italic=True); n.alignment = Alignment(wrap_text=True, vertical='center', indent=1)
ws.row_dimensions[2].height = 34

cab = ['Nº', 'ABA', 'AMBIENTE', 'GRUPO', 'ITEM', 'CÉLULA', 'ESPECIFICAÇÕES CONFLITANTES (versões e fontes)', 'DEFINIÇÃO DA ARQUITETURA']
for i, v in enumerate(cab, 1):
    c = ws.cell(4, i, v)
    c.font = Font(name='Aptos Narrow', size=11, bold=True); c.fill = PatternFill('solid', fgColor=CINZA)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BORDA
ws.row_dimensions[4].height = 30

for k, (aba, amb, grupo, item, ref, aba_real, texto) in enumerate(itens, 1):
    r = 4 + k
    vals = [k, aba, amb, grupo, item, ref, texto, '']
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v)
        c.font = Font(name='Aptos Narrow', size=11, bold=(i in (3, 6)))
        c.alignment = Alignment(horizontal='left' if i == 7 else 'center', vertical='center', wrap_text=True)
        c.border = BORDA
    g = ws.cell(r, 7); g.fill = PatternFill('solid', fgColor=VERM); g.font = Font(name='Aptos Narrow', size=11, color=VERM_TXT)
    f = ws.cell(r, 6)
    f.hyperlink = Hyperlink(ref=f.coordinate, location=f"'{aba_real}'!{ref}", tooltip=f'Ir para {aba} {ref}')
    f.font = Font(name='Aptos Narrow', size=11, bold=True, color='0563C1', underline='single')
    linhas = sum(max(1, -(-len(p) // 95)) for p in texto.split('\n'))
    ws.row_dimensions[r].height = max(30, min(409, linhas * 15 + 12))

ws.freeze_panes = 'A5'
ws.print_title_rows = '1:4'
ps = ws.page_setup; ps.paperSize = ws.PAPERSIZE_A4; ps.orientation = 'landscape'
ws.sheet_properties.pageSetUpPr.fitToPage = True; ps.fitToWidth = 1; ps.fitToHeight = 0
ws.print_options.horizontalCentered = True
m = ws.page_margins; m.left = m.right = 0.3; m.top = m.bottom = 0.4
ws.oddFooter.left.text = wb['TÉRREO '].oddFooter.left.text
wb.save(P)
print(len(itens), 'conflitos')
for it in itens: print(' ', it[0], it[4], it[1], '|', it[3])
