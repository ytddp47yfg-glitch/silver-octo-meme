"""Imagens dentro da própria aba da grade, com link nas células.

As fotos ficam à direita da grade (a partir da coluna AH, fora da área de
impressão), na mesma linha do ambiente. A célula que cita o código recebe a
linha "► VER IMAGEM (clique)" e um hyperlink interno para a célula da foto;
abaixo de cada foto há "◄ VOLTAR" para a célula de origem.
Não há aba separada de imagens. Rodar depois de formatacao.py e antes de
area_impressao.py (formatacao.py apaga as larguras das colunas extras).
"""
import sys, io, copy
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image

P = sys.argv[1] if len(sys.argv) > 1 else 'especificacoes/CDJ-GRE-ARQ-AC-R00.xlsx'
DIR = 'especificacoes/imagens/'
MARCA = '► VER IMAGEM (clique)'
COL0 = 34            # AH: primeira coluna de imagens (a grade vai até AF)
LARG = 72.0          # largura das colunas de imagem
BOX_W = 500          # px úteis para a foto
# (aba, célula que cita o item) -> (arquivo da imagem, legenda)
LINKS = [
    ('PILOTIS', 'B41', 'CDJ-IMG-01_ESPACO-TERMAS_P10.jpg', 'P10 - Piso elevado em mármore branco selecionado'),
    ('PILOTIS', 'H41', 'CDJ-IMG-02-04_ESPACO-TERMAS_PA49.jpg', 'PA49 - Mosaico Bisazza (painel e pilar)'),
    ('PILOTIS', 'U41', 'CDJ-IMG-03_ESPACO-TERMAS_D54.jpg', 'D54 - Chuveiro de embutir no teto Raindream, Roca'),
    ('PILOTIS', 'B55', 'CDJ-IMG-05_VESTIARIO-FEM-MASC-TERMAS_P10.jpg', 'P10 - Piso elevado em mármore branco selecionado'),
    ('PILOTIS', 'H55', 'CDJ-IMG-12_VESTIARIO-FEM-MASC-TERMAS_PA47.jpg', 'PA47 - Pintura com tinta mineral, cor 930, Terracor'),
    ('PILOTIS', 'J55', 'CDJ-IMG-07_VESTIARIO-FEM-MASC-TERMAS_D73.jpg', 'D73 - Bancada em mármore branco selecionado'),
    ('PILOTIS', 'K55', 'CDJ-IMG-06_VESTIARIO-FEM-MASC-TERMAS_D38.jpg', 'D38 - Espelho liso incolor 4mm'),
    ('PILOTIS', 'L55', 'CDJ-IMG-08_VESTIARIO-FEM-MASC-TERMAS_D72.jpg', 'D72 - Cuba de semi-encaixe slim, Deca'),
    ('PILOTIS', 'O55', 'CDJ-IMG-09_VESTIARIO-FEM-MASC-TERMAS_D69.jpg', 'D69 - Torneira com sensor Decalux, Deca'),
    ('PILOTIS', 'P55', 'CDJ-IMG-10_VESTIARIO-FEM-MASC-TERMAS_D22.jpg', 'D22 - Bacia Carrara / Panache e assento'),
    ('PILOTIS', 'S55', 'CDJ-IMG-11_VESTIARIO-FEM-MASC-TERMAS_D33.jpg', 'D33 - Monocomando para chuveiro Mix&Match, Docol'),
    ('PILOTIS', 'U55', 'CDJ-IMG-13_VESTIARIO-FEM-MASC-TERMAS_D70.jpg', 'D70 - Chuveiro 200 de parede Docol Eden'),
]

def img_bytes(path, max_w=1600):
    im = Image.open(path).convert('RGB')
    if im.width > max_w:
        im = im.resize((max_w, int(im.height * max_w / im.width)))
    b = io.BytesIO(); im.save(b, 'JPEG', quality=85)
    return b.getvalue(), im.size

wb = openpyxl.load_workbook(P)
# guarda os bytes das imagens já existentes (cabeçalhos) antes de mexer
for ws in wb.worksheets:
    for im in ws._images:
        b = im._data(); im._data = (lambda b=b: b)

proxima = {}
abas = {a for a, *_ in LINKS}
for aba in abas:
    ws = wb[aba]
    ws._images = [im for im in ws._images if im.anchor._from.col < COL0 - 1]   # tira fotos de rodadas anteriores
    for c in range(COL0, COL0 + 12):
        ws.column_dimensions[L(c)].width = LARG
for aba, ref, arq, leg in LINKS:
    ws = wb[aba]
    cel = ws[ref]; r = cel.row
    col = proxima.get((aba, r), COL0); proxima[(aba, r)] = col + 1
    alvo = f'{L(col)}{r}'
    # célula da grade -> foto
    v = str(cel.value or '')
    if MARCA not in v:
        cel.value = v.rstrip() + '\n\n' + MARCA
    cel.hyperlink = Hyperlink(ref=ref, location=f"'{aba}'!{alvo}", tooltip=('Ver imagem: ' + leg)[:255], display=None)
    f = copy.copy(cel.font); f.u = None; cel.font = f
    # célula da foto: legenda + voltar (texto embaixo, foto em cima)
    t = ws[alvo]
    t.value = f'{leg}\n◄ VOLTAR para {ref}'
    t.hyperlink = Hyperlink(ref=alvo, location=f"'{aba}'!{ref}", tooltip='Voltar para a grade')
    t.font = Font(name='Aptos Narrow', size=11, bold=True, color='243F2E')
    t.alignment = Alignment(horizontal='center', vertical='bottom', wrap_text=True)
    t.fill = PatternFill('solid', fgColor='FFFFFF')
    # foto ajustada à caixa (largura da coluna x altura da linha menos a legenda)
    data, (w, h) = img_bytes(DIR + arq)
    row_px = int((ws.row_dimensions[r].height or 60) * 96 / 72)
    box_h = max(80, row_px - 50)
    k = min(BOX_W / w, box_h / h)
    dw, dh = int(w * k), int(h * k)
    xi = XLImage(io.BytesIO(data)); xi._data = (lambda b=data: b)
    xi.width, xi.height = dw, dh
    col_px = int(LARG * 7 + 5)
    off_x = max(0, (col_px - dw) // 2)
    xi.anchor = OneCellAnchor(_from=AnchorMarker(col=col - 1, colOff=pixels_to_EMU(off_x), row=r - 1, rowOff=pixels_to_EMU(6)),
                              ext=XDRPositiveSize2D(pixels_to_EMU(dw), pixels_to_EMU(dh)))
    ws.add_image(xi)
    print(aba, ref, '->', alvo, arq, (dw, dh))
# título da área de imagens
for aba in abas:
    ws = wb[aba]
    h = ws.cell(11, COL0); h.value = 'IMAGENS DAS ESPECIFICAÇÕES (fora da área de impressão)'
    h.font = Font(name='Aptos Narrow', size=20, bold=True, color='FFFFFF'); h.fill = PatternFill('solid', fgColor='243F2E')
    h.alignment = Alignment(horizontal='left', vertical='center')
    rng = f'{L(COL0)}11:{L(COL0 + 9)}11'
    if rng not in [str(m) for m in ws.merged_cells.ranges]:
        ws.merge_cells(rng)

cr = wb['CONTROLE DE REVISÕES ']
for rr in range(12, 40):
    for c in range(1, 10):
        x = cr.cell(rr, c)
        if isinstance(x.value, str) and x.value.strip() in ('IMAGENS ESPECIFICAÇÕES', 'IMAGENS: LINK "VER IMAGEM" NAS CÉLULAS DA GRADE'):
            x.value = 'IMAGENS: AO LADO DA GRADE (LINK "VER IMAGEM")'
        if c == 4 and isinstance(x.value, str) and x.value.strip() == 'IMAGENS ESPECIFICAÇÕES':
            pass
for rr in range(24, 40):
    if str(cr.cell(rr, 4).value or '').strip() in ('IMAGENS ESPECIFICAÇÕES', 'IMAGENS: AO LADO DA GRADE (LINK "VER IMAGEM")'):
        cr.cell(rr, 4).value = 'IMAGENS ESPECIFICAÇÕES'
        cr.cell(rr, 5).value = ('Imagens das fichas TERMAS e VESTIÁRIOS - TERMAS da tabela de acabamentos v00. '
                                'Em 02/10/2026 passaram para a própria aba PILOTIS, à direita da grade (a partir da coluna AH); '
                                'o link "► VER IMAGEM (clique)" na célula leva à foto e "◄ VOLTAR" retorna.')
wb.save(P)
