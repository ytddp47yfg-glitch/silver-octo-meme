"""Prévia de impressão: cópia da xlsx com área de impressão fixa e sem hyperlinks.

O LibreOffice ignora o Print_Area dinâmico (OFFSET) e imprime as fotos da coluna AH;
também desenha célula com hyperlink numa linha só. Uso: previa_impressao.py <xlsx> <saida.xlsx>,
depois soffice --headless --convert-to pdf <saida.xlsx>.
"""
import openpyxl, sys
from openpyxl.utils import get_column_letter as L
wb=openpyxl.load_workbook(sys.argv[1])
for ws in wb.worksheets:
    for im in ws._images:
        b=im._data(); im._data=(lambda b=b:b)
cfg={'CONTROLE DE REVISÕES ':9,'TÉRREO ':32,'2° PAVTO':32,'PILOTIS':32,'APTO TIPO':32,'CONFLITOS':8,'PENDÊNCIAS':5}
for n,mc in cfg.items():
    ws=wb[n]; last=1
    for row in ws.iter_rows(max_col=mc):
        for c in row:
            if c.value not in (None,''): last=max(last,c.row)
    ws.print_area=f'A1:{L(mc)}{last}'
# o LibreOffice desenha célula com hyperlink numa linha só (sem quebra); na prévia tira os links
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.hyperlink is not None:
                c.hyperlink = None
wb.save(sys.argv[2])
