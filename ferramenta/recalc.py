#!/usr/bin/env python3
"""Recalcula todas as fórmulas de uma planilha .xlsx com o LibreOffice (sem interface) e conta células com erro.

Uso: python3 ferramenta/recalc.py planilha.xlsx [tempo_limite_s]
Saída (JSON): {"status": "ok" | "erro", "total_errors": N, "exemplos": [...], "msg": "..."}

Funciona em Windows, macOS e Linux. O LibreOffice é procurado em: variável SOFFICE, PATH (soffice / soffice.exe /
libreoffice) e nas pastas padrão de instalação. Usa um perfil próprio e temporário com uma macro
"calcular tudo e salvar", para não mexer no LibreOffice do usuário.
"""
import glob, json, os, shutil, subprocess, sys, tempfile
from pathlib import Path

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub RecalcularSalvar()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>"""
ERROS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:", "#VALOR!", "#NOME?", "#REF")


def soffice():
    cands = [os.environ.get("SOFFICE")] + [shutil.which(n) for n in ("soffice", "soffice.exe", "libreoffice")]
    cands += glob.glob(r"C:\Program Files*\LibreOffice*\program\soffice.exe")
    cands += ["/Applications/LibreOffice.app/Contents/MacOS/soffice", "/usr/bin/soffice", "/opt/libreoffice/program/soffice"]
    return next((c for c in cands if c and os.path.exists(c)), None)


def recalc(arq, limite=600):
    exe = soffice()
    if not exe:
        return {"status": "erro", "msg": "LibreOffice não encontrado (instale ou defina a variável SOFFICE)"}
    arq = str(Path(arq).resolve())
    with tempfile.TemporaryDirectory() as prof:
        uri = Path(prof).as_uri()
        subprocess.run([exe, "--headless", "--norestore", "--terminate_after_init", f"-env:UserInstallation={uri}"],
                       capture_output=True, timeout=120)
        mod = Path(prof) / "user" / "basic" / "Standard"
        if not mod.exists():
            return {"status": "erro", "msg": "o LibreOffice não criou o perfil temporário"}
        (mod / "Module1.xba").write_text(MACRO, encoding="utf-8")
        antes = os.stat(arq).st_mtime_ns
        subprocess.run([exe, "--headless", "--norestore", f"-env:UserInstallation={uri}",
                        "vnd.sun.star.script:Standard.Module1.RecalcularSalvar?language=Basic&location=application", arq],
                       capture_output=True, timeout=limite)
        if os.stat(arq).st_mtime_ns == antes:
            return {"status": "erro", "msg": "o LibreOffice não regravou a planilha (fórmulas NÃO recalculadas)"}
    from openpyxl import load_workbook
    wb = load_workbook(arq, data_only=True, read_only=True); ex = []; n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith(ERROS):
                    n += 1
                    if len(ex) < 20: ex.append(f"{ws.title}!{c.coordinate} {v}")
    return {"status": "ok", "total_errors": n, "exemplos": ex}


if __name__ == "__main__":
    r = recalc(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else 600)
    print(json.dumps(r, ensure_ascii=False, indent=1))
    sys.exit(0 if r["status"] == "ok" else 1)
